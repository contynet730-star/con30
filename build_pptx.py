"""
石神井南中学校 6/3 校内研修 講義スライド（30枚・約40分）
構成：参考PPTX「主体的に学習に取り組む態度の評価」（光が丘三中/髙橋指導主事）を踏襲
Part1：主体的・対話的で深い学び × 評価の在り方
Part2：「主体的に学習に取り組む態度」の評価
（評価のばらつきを防ぐ指導の工夫を Part2 内で扱う）

デザイン：練馬中PPTX テイスト
- フォント：メイリオ
- 大きな文字（28-40pt）
- 枠の色：水色
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu, Cm
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

WATER_BLUE = RGBColor(0x5B, 0xC0, 0xDE)
LIGHT_WATER = RGBColor(0xD6, 0xEE, 0xFA)
DEEP_WATER = RGBColor(0x2A, 0x96, 0xC2)
NAVY = RGBColor(0x1F, 0x3A, 0x6E)
ORANGE_KW = RGBColor(0xEB, 0x6C, 0x15)
RED = RGBColor(0xC0, 0x40, 0x40)
GREEN = RGBColor(0x07, 0xA9, 0x73)
YELLOW = RGBColor(0xFF, 0xE0, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
DARK_GRAY = RGBColor(0x40, 0x40, 0x40)

F_MAIN = "メイリオ"

prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)
SW = prs.slide_width
SH = prs.slide_height

def add_blank_slide():
    return prs.slides.add_slide(prs.slide_layouts[6])

def add_chapter_bar(slide, chapter_text):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Emu(720000))
    bar.fill.solid(); bar.fill.fore_color.rgb = DEEP_WATER
    bar.line.fill.background()
    tf = bar.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = Emu(300000); tf.margin_right = Emu(100000)
    tf.margin_top = Emu(20000); tf.margin_bottom = Emu(20000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = chapter_text
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = F_MAIN

def add_sub_heading(slide, text):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(830000), SW, Emu(700000))
    bar.fill.solid(); bar.fill.fore_color.rgb = LIGHT_WATER
    bar.line.fill.background()
    tf = bar.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = Emu(300000); tf.margin_right = Emu(200000)
    tf.margin_top = Emu(20000); tf.margin_bottom = Emu(20000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r.font.name = F_MAIN
    ul = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(1530000), SW, Emu(40000))
    ul.fill.solid(); ul.fill.fore_color.rgb = WATER_BLUE
    ul.line.fill.background()

def add_page_number(slide, num):
    tb = slide.shapes.add_textbox(Emu(11250000), Emu(6420000), Emu(800000), Emu(360000))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = 0; tf.margin_right = Emu(40000); tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = str(num)
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = DARK_GRAY
    r.font.name = F_MAIN

def add_source_line(slide, text):
    tb = slide.shapes.add_textbox(Emu(300000), Emu(6430000), Emu(10800000), Emu(360000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = Emu(50000); tf.margin_right = Emu(50000)
    tf.margin_top = Emu(10000); tf.margin_bottom = Emu(10000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "出典：" + text
    r.font.size = Pt(13)
    r.font.color.rgb = DARK_GRAY
    r.font.name = F_MAIN

def add_body_box(slide, left, top, width, height, lines,
                 default_size=28, default_color=BLACK, line_spacing=1.4,
                 align=PP_ALIGN.LEFT, font=F_MAIN, fill_color=None,
                 border_color=WATER_BLUE, border_width=3.0,
                 vertical_anchor=MSO_ANCHOR.TOP, default_bold=True):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    if fill_color:
        box.fill.solid(); box.fill.fore_color.rgb = fill_color
    else:
        box.fill.solid(); box.fill.fore_color.rgb = WHITE
    if border_color:
        box.line.color.rgb = border_color
        box.line.width = Pt(border_width)
    else:
        box.line.fill.background()
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = Emu(160000); tf.margin_right = Emu(160000)
    tf.margin_top = Emu(100000); tf.margin_bottom = Emu(100000)
    tf.vertical_anchor = vertical_anchor
    for i, item in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        if isinstance(item, list):
            for j, run_item in enumerate(item):
                if isinstance(run_item, tuple):
                    rt, ropts = run_item
                else:
                    rt, ropts = run_item, {}
                if j == 0:
                    p.alignment = ropts.get('align', align)
                    p.line_spacing = ropts.get('line_spacing', line_spacing)
                r = p.add_run()
                r.text = rt
                r.font.size = Pt(ropts.get('size', default_size))
                r.font.bold = ropts.get('bold', default_bold)
                r.font.color.rgb = ropts.get('color', default_color)
                r.font.name = ropts.get('font', font)
        else:
            if isinstance(item, tuple):
                text, opts = item
            else:
                text, opts = item, {}
            p.alignment = opts.get('align', align)
            p.line_spacing = opts.get('line_spacing', line_spacing)
            if 'space_before' in opts:
                p.space_before = Pt(opts['space_before'])
            r = p.add_run()
            r.text = text
            r.font.size = Pt(opts.get('size', default_size))
            r.font.bold = opts.get('bold', default_bold)
            r.font.color.rgb = opts.get('color', default_color)
            r.font.name = opts.get('font', font)

def add_plain_text(slide, left, top, width, height, lines,
                   default_size=24, default_color=NAVY, line_spacing=1.3,
                   align=PP_ALIGN.LEFT, font=F_MAIN,
                   vertical_anchor=MSO_ANCHOR.TOP, default_bold=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = Emu(50000); tf.margin_right = Emu(50000)
    tf.margin_top = Emu(30000); tf.margin_bottom = Emu(30000)
    tf.vertical_anchor = vertical_anchor
    for i, item in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        if isinstance(item, list):
            for j, run_item in enumerate(item):
                if isinstance(run_item, tuple):
                    rt, ropts = run_item
                else:
                    rt, ropts = run_item, {}
                if j == 0:
                    p.alignment = ropts.get('align', align)
                    p.line_spacing = ropts.get('line_spacing', line_spacing)
                r = p.add_run()
                r.text = rt
                r.font.size = Pt(ropts.get('size', default_size))
                r.font.bold = ropts.get('bold', default_bold)
                r.font.color.rgb = ropts.get('color', default_color)
                r.font.name = ropts.get('font', font)
        else:
            if isinstance(item, tuple):
                text, opts = item
            else:
                text, opts = item, {}
            p.alignment = opts.get('align', align)
            p.line_spacing = opts.get('line_spacing', line_spacing)
            r = p.add_run()
            r.text = text
            r.font.size = Pt(opts.get('size', default_size))
            r.font.bold = opts.get('bold', default_bold)
            r.font.color.rgb = opts.get('color', default_color)
            r.font.name = opts.get('font', font)

# ============================================================
# S1: 表紙
# ============================================================
s = add_blank_slide()
top_band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Emu(900000))
top_band.fill.solid(); top_band.fill.fore_color.rgb = DEEP_WATER
top_band.line.fill.background()
bot_band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(6300000), SW, Emu(558000))
bot_band.fill.solid(); bot_band.fill.fore_color.rgb = DEEP_WATER
bot_band.line.fill.background()
ln1 = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(910000), SW, Emu(40000))
ln1.fill.solid(); ln1.fill.fore_color.rgb = WATER_BLUE
ln1.line.fill.background()
ln2 = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(6250000), SW, Emu(40000))
ln2.fill.solid(); ln2.fill.fore_color.rgb = WATER_BLUE
ln2.line.fill.background()

add_plain_text(s, Emu(700000), Emu(180000), Emu(10800000), Emu(540000), [
    ("令和８年度　石神井南中学校　校内研修", {'size': 24, 'color': WHITE, 'bold': True, 'align': PP_ALIGN.LEFT}),
], default_size=24, vertical_anchor=MSO_ANCHOR.MIDDLE)

add_plain_text(s, Emu(400000), Emu(1450000), Emu(11400000), Emu(900000), [
    ("「主体的に学習に取り組む態度」の評価",
     {'size': 44, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=44, vertical_anchor=MSO_ANCHOR.MIDDLE)

add_body_box(s, Emu(1500000), Emu(2600000), Emu(9200000), Emu(900000), [
    ("〜「主体的・対話的で深い学び」と学習評価〜",
     {'size': 28, 'bold': True, 'color': DEEP_WATER, 'align': PP_ALIGN.CENTER}),
], default_size=28, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=4.0)

add_plain_text(s, Emu(400000), Emu(3800000), Emu(11400000), Emu(700000), [
    ("評価のばらつきを防ぐ、具体的な評価の在り方",
     {'size': 22, 'color': DARK_GRAY, 'align': PP_ALIGN.CENTER}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE)

add_plain_text(s, Emu(400000), Emu(5000000), Emu(11400000), Emu(440000), [
    ("令和８年６月３日（水）",
     {'size': 22, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE)
add_plain_text(s, Emu(400000), Emu(5470000), Emu(11400000), Emu(380000), [
    ("練馬区教育委員会　指導主事",
     {'size': 16, 'color': DARK_GRAY, 'align': PP_ALIGN.CENTER}),
], default_size=16, vertical_anchor=MSO_ANCHOR.MIDDLE)
add_plain_text(s, Emu(400000), Emu(5860000), Emu(11400000), Emu(420000), [
    ("紺多　章一郎",
     {'size': 24, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
], default_size=24, vertical_anchor=MSO_ANCHOR.MIDDLE)

# ============================================================
# S2: 本日の流れ
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "本日の流れ")
add_sub_heading(s, "講話 約40分 ／ グループワーク 約30分 ／ 講評 約10分")
add_page_number(s, 2)

agenda = [
    ("Part 1", "「主体的・対話的で深い学び」と評価の在り方", "約15分"),
    ("Part 2", "「主体的に学習に取り組む態度」の評価", "約20分"),
    ("まとめ", "評価のばらつきを防ぐ ─ 明日から動かす", "約5分"),
]
y = 1900000
for i, (part, title, time) in enumerate(agenda):
    yy = y + i * 1100000
    add_body_box(s, Emu(400000), Emu(yy), Emu(1900000), Emu(950000), [
        (part, {'size': 26, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
    ], default_size=26, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
    add_body_box(s, Emu(2400000), Emu(yy), Emu(7600000), Emu(950000), [
        (title, {'size': 24, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.LEFT}),
    ], default_size=24, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=2.5)
    add_body_box(s, Emu(10100000), Emu(yy), Emu(1700000), Emu(950000), [
        (time, {'size': 22, 'bold': True, 'color': DEEP_WATER, 'align': PP_ALIGN.CENTER}),
    ], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=2.5)

add_body_box(s, Emu(400000), Emu(5450000), Emu(11400000), Emu(800000), [
    [("講話のあと、", {'size': 22, 'color': BLACK}),
     ("各教科の評価資料を持ち寄ってグループ協議", {'size': 24, 'color': ORANGE_KW}),
     ("します", {'size': 22, 'color': BLACK})],
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=LIGHT_WATER, border_width=3.0)

# ============================================================
# S3: 校長依頼の確認
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "校長からの依頼")
add_sub_heading(s, "本日の研修で扱う3点")
add_page_number(s, 3)

items = [
    ("①", "「主体的・対話的で深い学び」を実現する", "評価の在り方"),
    ("②", "「主体的に学習に取り組む態度」", "の評価の実施"),
    ("③", "評価のばらつきが出ないようにする", "指導の工夫"),
]
y = 1900000
for i, (num, kw, tail) in enumerate(items):
    yy = y + i * 1300000
    add_body_box(s, Emu(400000), Emu(yy), Emu(1000000), Emu(1100000), [
        (num, {'size': 50, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
    ], default_size=50, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
    add_body_box(s, Emu(1500000), Emu(yy), Emu(10300000), Emu(1100000), [
        (kw, {'size': 24, 'color': ORANGE_KW}),
        (tail, {'size': 24, 'color': NAVY}),
    ], default_size=24, line_spacing=1.2, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0)

add_source_line(s, "石神井南中学校 校長依頼（指導主事派遣願）／木原校長メール（5/14）")

# ============================================================
# S4: ウォームアップ ─ 数字あてクイズ
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "ウォームアップ")
add_sub_heading(s, "この数字は何を表しているでしょう？")
add_page_number(s, 4)

add_body_box(s, Emu(400000), Emu(1900000), Emu(3700000), Emu(1500000), [
    ("47.2", {'size': 50, 'bold': True, 'color': DEEP_WATER, 'align': PP_ALIGN.CENTER}),
    ("30.5　／　22.3", {'size': 26, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=26, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=4.0)
add_plain_text(s, Emu(400000), Emu(3450000), Emu(3700000), Emu(400000), [
    ("A中学校", {'size': 22, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE)

add_body_box(s, Emu(4250000), Emu(1900000), Emu(3700000), Emu(1500000), [
    ("8.1", {'size': 50, 'bold': True, 'color': DEEP_WATER, 'align': PP_ALIGN.CENTER}),
    ("65.8　／　26.1", {'size': 26, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=26, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=4.0)
add_plain_text(s, Emu(4250000), Emu(3450000), Emu(3700000), Emu(400000), [
    ("B中学校", {'size': 22, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE)

add_body_box(s, Emu(8100000), Emu(1900000), Emu(3700000), Emu(1500000), [
    ("12.4", {'size': 50, 'bold': True, 'color': DEEP_WATER, 'align': PP_ALIGN.CENTER}),
    ("50.3　／　37.3", {'size': 26, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=26, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=4.0)
add_plain_text(s, Emu(8100000), Emu(3450000), Emu(3700000), Emu(400000), [
    ("C中学校", {'size': 22, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE)

add_body_box(s, Emu(400000), Emu(4150000), Emu(11400000), Emu(800000), [
    ("答え：区内3校・同一教科の「主体的に学習に取り組む態度」評定割合（A:B:C %）",
     {'size': 20, 'bold': True, 'color': ORANGE_KW, 'align': PP_ALIGN.CENTER}),
], default_size=20, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=LIGHT_WATER, border_width=3.0)

add_body_box(s, Emu(400000), Emu(5050000), Emu(11400000), Emu(1100000), [
    ("同じ教科でも、学校で大きく異なる現実",
     {'size': 26, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
    [("→ ", {'size': 22, 'color': NAVY}),
     ("「主体的に学習に取り組む態度」の評価", {'size': 22, 'color': ORANGE_KW}),
     ("は、なぜ揺れるのか？", {'size': 22, 'color': NAVY})],
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0, line_spacing=1.3)

# ============================================================
# S5: Part1 章扉
# ============================================================
s = add_blank_slide()
add_page_number(s, 5)
add_body_box(s, Emu(1000000), Emu(2000000), Emu(10200000), Emu(2900000), [
    ("Part 1", {'size': 36, 'bold': True, 'color': WATER_BLUE, 'align': PP_ALIGN.CENTER}),
    ("", {'size': 14}),
    ("「主体的・対話的で深い学び」", {'size': 38, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
    ("と評価の在り方", {'size': 38, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
    ("", {'size': 8}),
    ("─ 何ができるようになるか × どのように学ぶか ─",
     {'size': 18, 'bold': True, 'color': DEEP_WATER, 'align': PP_ALIGN.CENTER}),
], default_size=28, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=5.0, fill_color=LIGHT_WATER)

# ============================================================
# S6: 学習指導要領 ─ 何ができるようになるか（資質・能力3つの柱）
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 1　深い学び × 評価の在り方")
add_sub_heading(s, "何ができるようになるか ─ 資質・能力の3つの柱")
add_page_number(s, 6)

pillars = [
    ("①", "知識・技能", "何を理解しているか、何ができるか"),
    ("②", "思考力・判断力・表現力等", "理解していること・できることをどう使うか"),
    ("③", "学びに向かう力・人間性等", "どのように社会・世界と関わり、よりよい人生を送るか"),
]
y = 1850000
for i, (num, name, desc) in enumerate(pillars):
    yy = y + i * 1380000
    add_body_box(s, Emu(400000), Emu(yy), Emu(900000), Emu(1100000), [
        (num, {'size': 44, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
    ], default_size=44, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
    add_body_box(s, Emu(1400000), Emu(yy), Emu(10400000), Emu(1100000), [
        (name, {'size': 26, 'color': ORANGE_KW}),
        (desc, {'size': 20, 'color': NAVY}),
    ], default_size=22, line_spacing=1.2, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0)

add_source_line(s, "中学校学習指導要領 解説　総則編／文部科学省")

# ============================================================
# S7: どのように学ぶか ─ 主体的・対話的で深い学び
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 1　深い学び × 評価の在り方")
add_sub_heading(s, "どのように学ぶか ─ 主体的・対話的で深い学び")
add_page_number(s, 7)

three = [
    ("主体的な学び", "見通しをもち、粘り強く取り組み、振り返る"),
    ("対話的な学び", "互いの考えを比べ、共に考えを創り上げる"),
    ("深い学び", "知識・技能を活用・概念化し、新たなものを創る"),
]
x = 400000
w = 3750000
gap = 80000
for i, (name, desc) in enumerate(three):
    xx = x + i * (w + gap)
    add_body_box(s, Emu(xx), Emu(1900000), Emu(w), Emu(700000), [
        (name, {'size': 28, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
    ], default_size=28, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
    add_body_box(s, Emu(xx), Emu(2620000), Emu(w), Emu(1800000), [
        (desc, {'size': 20, 'color': NAVY, 'align': PP_ALIGN.LEFT}),
    ], default_size=20, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0, line_spacing=1.4)

add_body_box(s, Emu(400000), Emu(4600000), Emu(11400000), Emu(1600000), [
    [("授業改善の視点 → ", {'size': 22, 'color': BLACK}),
     ("生徒に見せたい「学びの姿」", {'size': 26, 'color': ORANGE_KW})],
    [("　 が ", {'size': 22, 'color': BLACK}),
     ("評価の対象", {'size': 26, 'color': ORANGE_KW}),
     ("（特に「主体的に取り組む態度」）", {'size': 22, 'color': BLACK})],
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=LIGHT_WATER, border_width=3.0, line_spacing=1.4)

add_source_line(s, "中学校学習指導要領 解説　総則編／文部科学省")

# ============================================================
# S8: 「主体的な学び」のイメージ
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 1　深い学び × 評価の在り方")
add_sub_heading(s, "「主体的な学び」のイメージ")
add_page_number(s, 8)

# 6つの要素を配置
items = [
    ("興味や関心\nを高める", 400000, 1900000),
    ("見通しを\nもつ", 4250000, 1900000),
    ("粘り強く\n取り組む", 8100000, 1900000),
    ("自分と結び\n付ける", 400000, 3550000),
    ("振り返って\n次へつなげる", 4250000, 3550000),
    ("学び続ける\n意志", 8100000, 3550000),
]
for (txt, xx, yy) in items:
    add_body_box(s, Emu(xx), Emu(yy), Emu(3700000), Emu(1500000), [
        (txt, {'size': 22, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
    ], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=LIGHT_WATER, border_width=3.0, line_spacing=1.2)

add_body_box(s, Emu(400000), Emu(5300000), Emu(11400000), Emu(900000), [
    [("これらが ", {'size': 22, 'color': BLACK}),
     ("単元の中で実際に見えるか", {'size': 24, 'color': ORANGE_KW}),
     ("を観る", {'size': 22, 'color': BLACK})],
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0)

add_source_line(s, "中学校学習指導要領 解説　総則編／文部科学省")

# ============================================================
# S9: 「対話的な学び」のイメージ
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 1　深い学び × 評価の在り方")
add_sub_heading(s, "「対話的な学び」のイメージ")
add_page_number(s, 9)

items = [
    ("互いの考え\nを比べる", 400000, 1900000),
    ("多様な手段\nで説明する", 4250000, 1900000),
    ("共に考えを\n創り上げる", 8100000, 1900000),
    ("多様な情報\nを収集する", 400000, 3550000),
    ("協働して\n課題を解決", 4250000, 3550000),
    ("先哲の考え\nを手掛かりに", 8100000, 3550000),
]
for (txt, xx, yy) in items:
    add_body_box(s, Emu(xx), Emu(yy), Emu(3700000), Emu(1500000), [
        (txt, {'size': 22, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
    ], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=LIGHT_WATER, border_width=3.0, line_spacing=1.2)

add_body_box(s, Emu(400000), Emu(5300000), Emu(11400000), Emu(900000), [
    [("Input ⇄ Output の往復で ", {'size': 22, 'color': BLACK}),
     ("自分の考えを更新する姿", {'size': 24, 'color': ORANGE_KW}),
     ("を観る", {'size': 22, 'color': BLACK})],
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0)

add_source_line(s, "中学校学習指導要領 解説　総則編／文部科学省")

# ============================================================
# S10: 「深い学び」のイメージ
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 1　深い学び × 評価の在り方")
add_sub_heading(s, "「深い学び」のイメージ")
add_page_number(s, 10)

items = [
    ("知識・技能を\n習得する", 400000, 1900000),
    ("知識・技能を\n活用する", 4250000, 1900000),
    ("知識・技能を\n概念化する", 8100000, 1900000),
    ("自分の考えを\n形成する", 400000, 3550000),
    ("思考して\n問い続ける", 4250000, 3550000),
    ("新たなものを\n創り上げる", 8100000, 3550000),
]
for (txt, xx, yy) in items:
    add_body_box(s, Emu(xx), Emu(yy), Emu(3700000), Emu(1500000), [
        (txt, {'size': 22, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
    ], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=LIGHT_WATER, border_width=3.0, line_spacing=1.2)

add_body_box(s, Emu(400000), Emu(5300000), Emu(11400000), Emu(900000), [
    [("「分かった」を超えて ", {'size': 22, 'color': BLACK}),
     ("「使える・問い続ける」姿", {'size': 24, 'color': ORANGE_KW}),
     ("を観る", {'size': 22, 'color': BLACK})],
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0)

add_source_line(s, "中学校学習指導要領 解説　総則編／文部科学省")

# ============================================================
# S11: 授業改善と評価の一体化
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 1　深い学び × 評価の在り方")
add_sub_heading(s, "授業改善 と 評価 は一体")
add_page_number(s, 11)

add_body_box(s, Emu(400000), Emu(1900000), Emu(11400000), Emu(900000), [
    ("「主体的・対話的で深い学び」の視点で授業を改善する",
     {'size': 26, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=26, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=LIGHT_WATER, border_width=3.0)

# 下向き矢印
arrow = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Emu(5596000), Emu(2900000), Emu(1000000), Emu(500000))
arrow.fill.solid(); arrow.fill.fore_color.rgb = ORANGE_KW
arrow.line.fill.background()

add_body_box(s, Emu(400000), Emu(3500000), Emu(11400000), Emu(900000), [
    ("授業で「目指す生徒の姿」を、教師が事前に描く",
     {'size': 26, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=26, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=LIGHT_WATER, border_width=3.0)

arrow2 = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Emu(5596000), Emu(4500000), Emu(1000000), Emu(500000))
arrow2.fill.solid(); arrow2.fill.fore_color.rgb = ORANGE_KW
arrow2.line.fill.background()

add_body_box(s, Emu(400000), Emu(5100000), Emu(11400000), Emu(1100000), [
    [("その「目指す姿」が ", {'size': 24, 'color': BLACK}),
     ("評価規準（Bと判断する姿）", {'size': 26, 'color': ORANGE_KW}),
     (" になる", {'size': 24, 'color': BLACK})],
], default_size=24, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=4.0)

# ============================================================
# S12: キーメッセージ①
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 1　深い学び × 評価の在り方")
add_sub_heading(s, "【キーメッセージ①】")
add_page_number(s, 12)

add_body_box(s, Emu(800000), Emu(2100000), Emu(10600000), Emu(1500000), [
    [("評価は", {'size': 36, 'color': NAVY}),
     ("「ためる」のでなく", {'size': 36, 'color': NAVY})],
    [("「", {'size': 42, 'color': NAVY}),
     ("使う", {'size': 50, 'color': ORANGE_KW}),
     ("」もの", {'size': 42, 'color': NAVY})],
], default_size=40, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=5.0,
   fill_color=LIGHT_WATER, align=PP_ALIGN.CENTER)

add_body_box(s, Emu(800000), Emu(4000000), Emu(10600000), Emu(1800000), [
    [("授業で「目指す姿」を描き、", {'size': 24, 'color': BLACK, 'align': PP_ALIGN.CENTER})],
    [("その姿が見えたかを", {'size': 24, 'color': BLACK, 'align': PP_ALIGN.CENTER}),
     ("評価", {'size': 28, 'color': ORANGE_KW, 'align': PP_ALIGN.CENTER}),
     ("する", {'size': 24, 'color': BLACK, 'align': PP_ALIGN.CENTER})],
    ("評価結果を、次の授業に「使う」", {'size': 22, 'color': DARK_GRAY, 'align': PP_ALIGN.CENTER}),
], default_size=24, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0, line_spacing=1.3)

# ============================================================
# S13: Part2 章扉
# ============================================================
s = add_blank_slide()
add_page_number(s, 13)
add_body_box(s, Emu(1000000), Emu(2000000), Emu(10200000), Emu(2900000), [
    ("Part 2", {'size': 36, 'bold': True, 'color': WATER_BLUE, 'align': PP_ALIGN.CENTER}),
    ("", {'size': 14}),
    ("「主体的に学習に取り組む態度」", {'size': 36, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
    ("の評価", {'size': 36, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
    ("", {'size': 8}),
    ("─ どう見取り、どうばらつきを防ぐか ─",
     {'size': 20, 'bold': True, 'color': DEEP_WATER, 'align': PP_ALIGN.CENTER}),
], default_size=28, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=5.0, fill_color=LIGHT_WATER)

# ============================================================
# S14: 学習評価の現状における課題
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "学習評価の現状における課題")
add_page_number(s, 14)

issues = [
    ("学習改善への活用が弱い", "「ためた評価」が次の授業に生きていない"),
    ("生徒・保護者からの誤解", "「挙手」「忘れ物」などで決まると思われている"),
    ("教員の負担", "観点が多すぎ、毎時間記録に追われる"),
    ("正しい理解が共有されていない", "「主体的に取り組む態度」の見取り方が揺れる"),
]
y = 1900000
for i, (kw, desc) in enumerate(issues):
    yy = y + i * 1050000
    add_body_box(s, Emu(400000), Emu(yy), Emu(11400000), Emu(950000), [
        [(kw, {'size': 26, 'color': ORANGE_KW}),
         ("　─　", {'size': 22, 'color': DARK_GRAY}),
         (desc, {'size': 20, 'color': NAVY})],
    ], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0)

add_body_box(s, Emu(400000), Emu(6080000), Emu(11400000), Emu(240000), [
    ("何のために「学習評価」を行うのか、改めて意識することが重要",
     {'size': 14, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=14, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=YELLOW, border_color=ORANGE_KW, border_width=2.0)

# ============================================================
# S15: 学習評価の改善の基本方針
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "学習評価の改善の基本方針")
add_page_number(s, 15)

basics = [
    ("①", "生徒の学習改善", "につながるものにしていく"),
    ("②", "教師の指導改善", "につながるものにしていく"),
    ("③", "必要性・妥当性", "が認められないものは見直していく"),
]
y = 1900000
for i, (num, kw, tail) in enumerate(basics):
    yy = y + i * 1300000
    add_body_box(s, Emu(400000), Emu(yy), Emu(900000), Emu(1100000), [
        (num, {'size': 44, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
    ], default_size=44, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
    add_body_box(s, Emu(1400000), Emu(yy), Emu(10400000), Emu(1100000), [
        [(kw, {'size': 28, 'color': ORANGE_KW}),
         (tail, {'size': 24, 'color': NAVY})],
    ], default_size=24, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0)

add_source_line(s, "平成31年１月 中央教育審議会「児童生徒の学習評価の在り方について（報告）」")

# ============================================================
# S16: 各教科の評価の基本構造（3観点）
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "各教科における評価の基本構造")
add_page_number(s, 16)

views = [
    ("知識・技能", "何を知っているか・できるか", DEEP_WATER),
    ("思考・判断・表現", "知っていること・できることをどう使うか", WATER_BLUE),
    ("主体的に学習に取り組む態度", "学びをどう自己調整しているか", ORANGE_KW),
]
y = 1900000
for i, (name, desc, col) in enumerate(views):
    yy = y + i * 1300000
    add_body_box(s, Emu(400000), Emu(yy), Emu(5000000), Emu(1100000), [
        (name, {'size': 26, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
    ], default_size=26, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=col, border_color=col, border_width=2.0)
    add_body_box(s, Emu(5500000), Emu(yy), Emu(6300000), Emu(1100000), [
        (desc, {'size': 22, 'color': NAVY, 'align': PP_ALIGN.LEFT}),
    ], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0)

add_source_line(s, "国立教育政策研究所「学習評価の在り方ハンドブック（中学校編）」")

# ============================================================
# S17: 「知識・技能」の評価
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "「知識・技能」の評価")
add_page_number(s, 17)

add_body_box(s, Emu(400000), Emu(1900000), Emu(11400000), Emu(1500000), [
    [("既有の知識・技能と", {'size': 24, 'color': BLACK}),
     ("関連付け・活用", {'size': 26, 'color': ORANGE_KW}),
     ("する中で", {'size': 24, 'color': BLACK})],
    [("　", {'size': 24}),
     ("概念として理解し、技能を習得", {'size': 26, 'color': ORANGE_KW}),
     ("しているか", {'size': 24, 'color': BLACK})],
], default_size=24, line_spacing=1.3, border_width=3.0, fill_color=LIGHT_WATER)

add_body_box(s, Emu(400000), Emu(3550000), Emu(11400000), Emu(640000), [
    ("＜評価の工夫例＞", {'size': 24, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.LEFT}),
], default_size=24, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)

add_body_box(s, Emu(400000), Emu(4190000), Emu(11400000), Emu(2000000), [
    ("○ ペーパーテスト", {'size': 24, 'color': NAVY}),
    ("", {'size': 6}),
    ("○ 実際に知識や技能を用いる場面を設けた学習活動", {'size': 24, 'color': NAVY}),
    ("", {'size': 6}),
    ("　 （実験・実演・作品制作 等）", {'size': 20, 'color': DARK_GRAY}),
], default_size=24, line_spacing=1.3, border_width=3.0)

add_source_line(s, "国立教育政策研究所「学習評価の在り方ハンドブック（中学校編）」")

# ============================================================
# S18: 「思考・判断・表現」の評価
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "「思考・判断・表現」の評価")
add_page_number(s, 18)

add_body_box(s, Emu(400000), Emu(1900000), Emu(11400000), Emu(1500000), [
    [("知識・技能を", {'size': 24, 'color': BLACK}),
     ("活用して課題を解決", {'size': 26, 'color': ORANGE_KW}),
     ("するために必要な", {'size': 24, 'color': BLACK})],
    [("　", {'size': 24}),
     ("思考力・判断力・表現力", {'size': 26, 'color': ORANGE_KW}),
     ("を身に付けているか", {'size': 24, 'color': BLACK})],
], default_size=24, line_spacing=1.3, border_width=3.0, fill_color=LIGHT_WATER)

add_body_box(s, Emu(400000), Emu(3550000), Emu(11400000), Emu(640000), [
    ("＜評価の工夫例＞", {'size': 24, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.LEFT}),
], default_size=24, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)

add_body_box(s, Emu(400000), Emu(4190000), Emu(11400000), Emu(2000000), [
    ("○ 論述やレポートの作成・発表", {'size': 22, 'color': NAVY}),
    ("○ グループでの話合い", {'size': 22, 'color': NAVY}),
    ("○ 作品の制作や表現等の多様な活動", {'size': 22, 'color': NAVY}),
    ("○ ポートフォリオの活用", {'size': 22, 'color': NAVY}),
], default_size=22, line_spacing=1.3, border_width=3.0)

add_source_line(s, "国立教育政策研究所「学習評価の在り方ハンドブック（中学校編）」")

# ============================================================
# S19: 「主体的に学習に取り組む態度」の評価 ─ 2つの側面
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "「主体的に学習に取り組む態度」── 2つの側面")
add_page_number(s, 19)

add_body_box(s, Emu(400000), Emu(1900000), Emu(5600000), Emu(700000), [
    ("① 粘り強い取組", {'size': 26, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
], default_size=26, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
add_body_box(s, Emu(400000), Emu(2600000), Emu(5600000), Emu(2000000), [
    ("自らの学習を", {'size': 22, 'color': BLACK}),
    [("　", {'size': 22}),
     ("粘り強く調整", {'size': 26, 'color': ORANGE_KW})],
    ("しようとしている姿", {'size': 22, 'color': BLACK}),
], default_size=22, line_spacing=1.3, border_width=3.0)

add_body_box(s, Emu(6200000), Emu(1900000), Emu(5600000), Emu(700000), [
    ("② 自己調整", {'size': 26, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
], default_size=26, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=WATER_BLUE, border_color=WATER_BLUE, border_width=2.0)
add_body_box(s, Emu(6200000), Emu(2600000), Emu(5600000), Emu(2000000), [
    ("自らの学習状況を", {'size': 22, 'color': BLACK}),
    [("　", {'size': 22}),
     ("把握し、調整", {'size': 26, 'color': ORANGE_KW})],
    ("しようとしている姿", {'size': 22, 'color': BLACK}),
], default_size=22, line_spacing=1.3, border_width=3.0)

add_body_box(s, Emu(400000), Emu(4800000), Emu(11400000), Emu(1400000), [
    [("①と②は", {'size': 22, 'color': BLACK}),
     ("別々ではなく相互に関わり合う", {'size': 24, 'color': ORANGE_KW})],
    ("①がＡで②がＣ という姿は 一般的ではない",
     {'size': 22, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=LIGHT_WATER, border_width=4.0, line_spacing=1.3)

add_source_line(s, "国立教育政策研究所「指導と評価の一体化のための参考資料」")

# ============================================================
# S20: 「主体的に学習に取り組む態度」の評価の工夫例
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "「主体的に取り組む態度」── 評価の工夫例")
add_page_number(s, 20)

add_body_box(s, Emu(400000), Emu(1900000), Emu(11400000), Emu(640000), [
    ("＜評価の工夫例＞", {'size': 24, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.LEFT}),
], default_size=24, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)

works = [
    ("○ ノートやレポート等における記述", "授業の前後・単元末の振り返り"),
    ("○ 授業中の発言・行動観察", "教師が見取りたい姿に対する具体的記録"),
    ("○ 生徒による自己評価・相互評価", "自己調整の見取り材料として活用"),
]
y = 2640000
for i, (kw, sub) in enumerate(works):
    yy = y + i * 1100000
    add_body_box(s, Emu(400000), Emu(yy), Emu(11400000), Emu(950000), [
        (kw, {'size': 24, 'color': NAVY}),
        (sub, {'size': 18, 'color': DARK_GRAY}),
    ], default_size=22, line_spacing=1.2, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0)

add_body_box(s, Emu(400000), Emu(6020000), Emu(11400000), Emu(300000), [
    ("単一の手段でなく、複数を組み合わせて見取る",
     {'size': 14, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=14, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=YELLOW, border_color=ORANGE_KW, border_width=2.0)

add_source_line(s, "国立教育政策研究所「指導と評価の一体化のための参考資料」")

# ============================================================
# S21: 行動観察の工夫（数学の例）
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "行動観察の工夫例 ─ 数学・問題解決の場面")
add_page_number(s, 21)

add_body_box(s, Emu(400000), Emu(1900000), Emu(5600000), Emu(640000), [
    ("①粘り強く取り組む姿", {'size': 22, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
add_body_box(s, Emu(400000), Emu(2540000), Emu(5600000), Emu(2700000), [
    ("○ 前時までのノートを", {'size': 18, 'color': NAVY}),
    ("　 見返している", {'size': 18, 'color': NAVY}),
    ("", {'size': 6}),
    ("○ 別の方法での解決を試みる", {'size': 18, 'color': NAVY}),
    ("", {'size': 6}),
    ("○ 友達に説明することを念頭に置いて、", {'size': 18, 'color': NAVY}),
    ("　 着想や説明を書き加えていく", {'size': 18, 'color': NAVY}),
], default_size=18, line_spacing=1.3, border_width=3.0)

add_body_box(s, Emu(6200000), Emu(1900000), Emu(5600000), Emu(640000), [
    ("②自己調整している姿", {'size': 22, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=WATER_BLUE, border_color=WATER_BLUE, border_width=2.0)
add_body_box(s, Emu(6200000), Emu(2540000), Emu(5600000), Emu(2700000), [
    ("○ 解決方法を振り返り、修正したり", {'size': 18, 'color': NAVY}),
    ("　 別の方法を考えたりする", {'size': 18, 'color': NAVY}),
    ("", {'size': 6}),
    ("○ 式に言葉や図での説明を書き加える", {'size': 18, 'color': NAVY}),
    ("", {'size': 6}),
    ("○ 解決方法を自己評価し、", {'size': 18, 'color': NAVY}),
    ("　 更によいものを求めようとする", {'size': 18, 'color': NAVY}),
], default_size=18, line_spacing=1.3, border_width=3.0)

add_body_box(s, Emu(400000), Emu(5400000), Emu(11400000), Emu(800000), [
    [("ポイント：", {'size': 20, 'color': BLACK}),
     ("生徒の学びの姿をイメージした学習活動", {'size': 22, 'color': ORANGE_KW}),
     ("をつくる", {'size': 20, 'color': BLACK})],
], default_size=20, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=LIGHT_WATER, border_width=3.0)

add_source_line(s, "国立教育政策研究所「指導と評価の一体化のための参考資料」（数学）")

# ============================================================
# S22: 自己評価・相互評価の工夫（評価の発問）
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "評価につながる発問の工夫")
add_page_number(s, 22)

prompts = [
    ("①", "理解の状況を振り返る発問", "授業の最初と比べて、考え方が変わったところは？"),
    ("②", "他者との協働で考えを相対化", "友達の意見を聞いて、自分の考えを見直してみよう"),
    ("③", "目標達成状況の振り返り", "今日の学習の目標は達成できたかな？"),
]
y = 1900000
for i, (num, name, ex) in enumerate(prompts):
    yy = y + i * 1300000
    add_body_box(s, Emu(400000), Emu(yy), Emu(900000), Emu(1150000), [
        (num, {'size': 44, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
    ], default_size=44, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
    add_body_box(s, Emu(1400000), Emu(yy), Emu(10400000), Emu(1150000), [
        (name, {'size': 22, 'color': ORANGE_KW}),
        ("発問例：「" + ex + "」", {'size': 18, 'color': NAVY}),
    ], default_size=20, line_spacing=1.2, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0)

add_source_line(s, "国立教育政策研究所「指導と評価の一体化のための参考資料」")

# ============================================================
# S23: キーメッセージ② 「指導してから評価する」
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "【キーメッセージ②】《本日の最重要メッセージ》")
add_page_number(s, 23)

add_body_box(s, Emu(400000), Emu(1900000), Emu(11400000), Emu(700000), [
    ("「主体的に学習に取り組む態度」は",
     {'size': 26, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
], default_size=26, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)

add_body_box(s, Emu(400000), Emu(2650000), Emu(11400000), Emu(900000), [
    ("“指導してから評価する”",
     {'size': 38, 'bold': True, 'color': ORANGE_KW, 'align': PP_ALIGN.CENTER}),
], default_size=38, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=LIGHT_WATER, border_width=5.0)

add_body_box(s, Emu(400000), Emu(3800000), Emu(900000), Emu(900000), [
    ("×", {'size': 44, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
], default_size=44, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=RED, border_color=RED, border_width=2.0)
add_body_box(s, Emu(1400000), Emu(3800000), Emu(10400000), Emu(900000), [
    [("「個人任せ」で見取ろうとする →", {'size': 20, 'color': NAVY}),
     ("　育たない・評価できない", {'size': 22, 'color': RED})],
], default_size=20, vertical_anchor=MSO_ANCHOR.MIDDLE, border_color=RED, border_width=3.0)

add_body_box(s, Emu(400000), Emu(4800000), Emu(900000), Emu(900000), [
    ("○", {'size': 44, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
], default_size=44, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=GREEN, border_color=GREEN, border_width=2.0)
add_body_box(s, Emu(1400000), Emu(4800000), Emu(10400000), Emu(900000), [
    [("「学び方」を指導してから見取る →", {'size': 20, 'color': NAVY}),
     ("　育つ・評価できる", {'size': 22, 'color': GREEN})],
], default_size=20, vertical_anchor=MSO_ANCHOR.MIDDLE, border_color=GREEN, border_width=3.0)

add_plain_text(s, Emu(400000), Emu(5850000), Emu(11400000), Emu(380000), [
    ("見取る前に、「学び方」を授業の中で教える",
     {'size': 18, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=18, vertical_anchor=MSO_ANCHOR.MIDDLE)

add_source_line(s, "国立教育政策研究所「指導と評価の一体化のための参考資料」")

# ============================================================
# S24: 9つの働きかけ（学び方を指導する）
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "「学び方」を指導する 9つの働きかけ")
add_page_number(s, 24)

works = [
    "1. 学習目標の共有",
    "2. 既習との関連付け",
    "3. 見通しを持たせる",
    "4. 振り返りの場面設定",
    "5. 自己評価の機会",
    "6. ペア・グループ対話",
    "7. 学習の選択肢提示",
    "8. ICTでの可視化",
    "9. 教師のモデリング",
]
y_start = 1900000
x_start = 400000
w = 3750000
h = 850000
gx = 80000
gy = 130000
for i, w_text in enumerate(works):
    row, col = divmod(i, 3)
    xx = x_start + col * (w + gx)
    yy = y_start + row * (h + gy)
    add_body_box(s, Emu(xx), Emu(yy), Emu(w), Emu(h), [
        (w_text, {'size': 22, 'color': NAVY, 'align': PP_ALIGN.LEFT}),
    ], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=LIGHT_WATER, border_width=3.0)

add_body_box(s, Emu(400000), Emu(5700000), Emu(11400000), Emu(550000), [
    ("1つ選んで、次の単元から実践する",
     {'size': 22, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=YELLOW, border_color=ORANGE_KW, border_width=3.0)

add_source_line(s, "国立教育政策研究所（2020）参考資料より整理")

# ============================================================
# S25: 評価のばらつきを防ぐ（妥当性×信頼性）
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "評価のばらつきを防ぐ ─ 妥当性 × 信頼性")
add_page_number(s, 25)

add_body_box(s, Emu(400000), Emu(1900000), Emu(5600000), Emu(700000), [
    ("妥当性", {'size': 26, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
], default_size=26, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
add_body_box(s, Emu(400000), Emu(2600000), Emu(5600000), Emu(2400000), [
    ("資質・能力を", {'size': 20, 'color': BLACK}),
    [("　", {'size': 20}),
     ("適切に反映", {'size': 24, 'color': ORANGE_KW})],
    ("していること", {'size': 20, 'color': BLACK}),
    ("", {'size': 8}),
    ("（「測りたい力」と「方法」", {'size': 18, 'color': DARK_GRAY}),
    ("　 が一致している）", {'size': 18, 'color': DARK_GRAY}),
], default_size=20, line_spacing=1.3, border_width=3.0)

add_body_box(s, Emu(6200000), Emu(1900000), Emu(5600000), Emu(700000), [
    ("信頼性", {'size': 26, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
], default_size=26, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=WATER_BLUE, border_color=WATER_BLUE, border_width=2.0)
add_body_box(s, Emu(6200000), Emu(2600000), Emu(5600000), Emu(2400000), [
    [("誰が評価しても", {'size': 20, 'color': BLACK})],
    [("　", {'size': 20}),
     ("同じ結果になる", {'size': 24, 'color': ORANGE_KW})],
    ("こと", {'size': 20, 'color': BLACK}),
    ("", {'size': 8}),
    ("（教師の主観だけに", {'size': 18, 'color': DARK_GRAY}),
    ("　 流されない）", {'size': 18, 'color': DARK_GRAY}),
], default_size=20, line_spacing=1.3, border_width=3.0)

add_body_box(s, Emu(400000), Emu(5200000), Emu(11400000), Emu(1000000), [
    [("妥当性は", {'size': 24, 'color': BLACK}),
     ("個人の授業設計", {'size': 26, 'color': ORANGE_KW}),
     ("で、", {'size': 24, 'color': BLACK}),
     ("信頼性は", {'size': 24, 'color': BLACK}),
     ("組織で揃える", {'size': 26, 'color': ORANGE_KW})],
], default_size=24, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=LIGHT_WATER, border_width=4.0)

add_source_line(s, "令和２年９月 東京都教育委員会「指導と評価の一体化を目指して」")

# ============================================================
# S26: 教科会で揃える3つ（ばらつき対策）
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "教科会で揃える3つ（ばらつき対策）")
add_page_number(s, 26)

items = [
    ("①", "Bの姿", "観点別評価で「B」と判断する具体的な姿"),
    ("②", "評価場面・方法", "いつ・どの活動で・どんな方法で見取るか"),
    ("③", "総括の方法", "観点別評価から評定への総括ルール"),
]
y = 1900000
for i, (num, kw, desc) in enumerate(items):
    yy = y + i * 1300000
    add_body_box(s, Emu(400000), Emu(yy), Emu(900000), Emu(1100000), [
        (num, {'size': 44, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
    ], default_size=44, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
    add_body_box(s, Emu(1400000), Emu(yy), Emu(10400000), Emu(1100000), [
        (kw, {'size': 26, 'color': ORANGE_KW}),
        (desc, {'size': 20, 'color': NAVY}),
    ], default_size=22, line_spacing=1.2, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0)

add_source_line(s, "東京都教育委員会「指導と評価の一体化を目指して」")

# ============================================================
# S27: Bの姿の具体化（数学の例）
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "「Bの姿」を具体化する ─ 数学「二次方程式」の例")
add_page_number(s, 27)

add_body_box(s, Emu(400000), Emu(1900000), Emu(11400000), Emu(640000), [
    ("学習シートに「気を付けるポイント」を記述",
     {'size': 22, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.LEFT}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)

add_body_box(s, Emu(400000), Emu(2600000), Emu(5600000), Emu(640000), [
    ("▼ おおむね満足（B）", {'size': 24, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
], default_size=24, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=WATER_BLUE, border_color=WATER_BLUE, border_width=2.0)
add_body_box(s, Emu(400000), Emu(3240000), Emu(5600000), Emu(2150000), [
    [("ポイントが", {'size': 22, 'color': BLACK}),
     ("書かれている", {'size': 24, 'color': ORANGE_KW})],
    ("", {'size': 8}),
    ("例：別の文字に置き換えて", {'size': 18, 'color': DARK_GRAY}),
    ("　 解く方法を使えるようにする", {'size': 18, 'color': DARK_GRAY}),
], default_size=20, line_spacing=1.3, border_width=3.0)

add_body_box(s, Emu(6200000), Emu(2600000), Emu(5600000), Emu(640000), [
    ("▼ 十分満足（A）", {'size': 24, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
], default_size=24, vertical_anchor=MSO_ANCHOR.MIDDLE,
   fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
add_body_box(s, Emu(6200000), Emu(3240000), Emu(5600000), Emu(2150000), [
    [("ポイント＋", {'size': 22, 'color': BLACK}),
     ("理由", {'size': 24, 'color': ORANGE_KW}),
     ("が書かれている", {'size': 22, 'color': BLACK})],
    ("", {'size': 8}),
    ("例：式を見ずに展開すると時間も", {'size': 18, 'color': DARK_GRAY}),
    ("　 間違いも増えるから、", {'size': 18, 'color': DARK_GRAY}),
    ("　 置き換えを使えるようにする", {'size': 18, 'color': DARK_GRAY}),
], default_size=20, line_spacing=1.3, border_width=3.0)

add_body_box(s, Emu(400000), Emu(5500000), Emu(11400000), Emu(700000), [
    ("同じ気付きでも「理由まで書けるか」でA／Bが分かれる",
     {'size': 22, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=LIGHT_WATER, border_width=4.0)

add_source_line(s, "国立教育政策研究所「指導と評価の一体化のための参考資料」（数学）")

# ============================================================
# S28: 単元計画に「記録欄」を1列足す
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "Part 2　「主体的に取り組む態度」の評価")
add_sub_heading(s, "単元計画に「記録」欄を1列足す")
add_page_number(s, 28)

add_body_box(s, Emu(400000), Emu(1900000), Emu(11400000), Emu(800000), [
    ("「いつ／何で／どう見取るか」を可視化",
     {'size': 26, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=26, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=LIGHT_WATER, border_width=4.0)

headers = ["時", "学習内容", "観点", "評価方法・場面"]
header_w = [1100000, 4500000, 2000000, 3800000]
x = 400000
y = 2900000
for i, hh in enumerate(headers):
    add_body_box(s, Emu(x), Emu(y), Emu(header_w[i]), Emu(600000), [
        (hh, {'size': 22, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
    ], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=1.5)
    x += header_w[i]

rows = [
    ("3-4", "気付きの記述", "態度", "学習シート"),
    ("7", "ペーパーテスト", "知技", "小テスト"),
]
y += 650000
for r in rows:
    x = 400000
    for i, txt in enumerate(r):
        add_body_box(s, Emu(x), Emu(y), Emu(header_w[i]), Emu(700000), [
            (txt, {'size': 22, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
        ], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=1.5)
        x += header_w[i]
    y += 730000

add_body_box(s, Emu(400000), Emu(5550000), Emu(11400000), Emu(640000), [
    ("全部記録しなくていい ─ 単元の節目で1観点ずつ",
     {'size': 22, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=YELLOW, border_color=ORANGE_KW, border_width=3.0)

add_source_line(s, "国立教育政策研究所「指導と評価の一体化のための参考資料」より作成")

# ============================================================
# S29: 明日からできること
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "まとめ ─ 明日から動かす")
add_sub_heading(s, "明日からできること ── 1つ選んでください")
add_page_number(s, 29)

choices = [
    "□ 次の単元の指導計画に「記録」欄を1列足す",
    "□ 教科会で「Bの姿」を1つ言語化して揃える",
    "□ 9つの働きかけから1つ選んで授業に入れる",
    "□ 振り返りシート／自己評価の機会を1回入れる",
]
y = 1900000
for i, txt in enumerate(choices):
    yy = y + i * 970000
    add_body_box(s, Emu(400000), Emu(yy), Emu(11400000), Emu(820000), [
        (txt, {'size': 24, 'color': NAVY}),
    ], default_size=24, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=LIGHT_WATER, border_width=3.0)

add_body_box(s, Emu(400000), Emu(5800000), Emu(11400000), Emu(500000), [
    ("やることを「絞る」から、続く",
     {'size': 22, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.CENTER}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=YELLOW, border_color=ORANGE_KW, border_width=3.0)

# ============================================================
# S30: 3つのキーメッセージ／協議会への橋渡し
# ============================================================
s = add_blank_slide()
add_chapter_bar(s, "まとめ ─ 本日のキーメッセージ")
add_sub_heading(s, "3つのキーメッセージ × 協議会へ")
add_page_number(s, 30)

msgs = [
    ("①", "評価は「ためる」より「使う」もの"),
    ("②", "「指導してから評価する」── 学び方を授業で教える"),
    ("③", "妥当性は個人で、信頼性は組織で"),
]
y = 1900000
for i, (num, msg) in enumerate(msgs):
    yy = y + i * 850000
    add_body_box(s, Emu(400000), Emu(yy), Emu(1000000), Emu(750000), [
        (num, {'size': 36, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}),
    ], default_size=36, vertical_anchor=MSO_ANCHOR.MIDDLE,
       fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
    add_body_box(s, Emu(1500000), Emu(yy), Emu(10300000), Emu(750000), [
        (msg, {'size': 22, 'color': NAVY}),
    ], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, border_width=3.0)

add_body_box(s, Emu(400000), Emu(4650000), Emu(11400000), Emu(640000), [
    ("▶  このあとの研究協議会では…",
     {'size': 22, 'bold': True, 'color': NAVY, 'align': PP_ALIGN.LEFT}),
], default_size=22, vertical_anchor=MSO_ANCHOR.MIDDLE, fill_color=YELLOW, border_color=ORANGE_KW, border_width=3.0)

add_body_box(s, Emu(400000), Emu(5350000), Emu(11400000), Emu(950000), [
    [("教科会単位で ", {'size': 20, 'color': BLACK}),
     ("各教科の評価資料を持ち寄り", {'size': 22, 'color': ORANGE_KW}),
     ("、", {'size': 20, 'color': BLACK})],
    [("　", {'size': 20}),
     ("「Bと判断する姿」を1つ言語化", {'size': 22, 'color': ORANGE_KW}),
     (" してみてください", {'size': 20, 'color': BLACK})],
], default_size=20, line_spacing=1.3, border_width=3.0)

# ========== 保存 ==========
import os
os.makedirs("/tmp/output", exist_ok=True)
prs.save("/tmp/output/講義スライド_石神井南中_20260603.pptx")
print("OK: PPTX created")
print(f"Slides: {len(prs.slides)}")
