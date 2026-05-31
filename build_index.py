# -*- coding: utf-8 -*-
"""
最終納品物一覧（インデックス文書）
6/3 研修当日に使う4ファイルを中心に整理
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
for sec in doc.sections:
    sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
style = doc.styles['Normal']
style.font.name = 'メイリオ'; style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'メイリオ')

NAVY = RGBColor(0x1F, 0x3A, 0x6E)
ORANGE = RGBColor(0xEB, 0x6C, 0x15)
GREEN = RGBColor(0x07, 0xA9, 0x73)
GRAY = RGBColor(0x40, 0x40, 0x40)
DEEP_WATER = RGBColor(0x2A, 0x96, 0xC2)
YELLOW = RGBColor(0xFF, 0xE0, 0x66)

def H1(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = NAVY
    r.font.name='メイリオ'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'メイリオ')

def H2(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = NAVY
    r.font.name='メイリオ'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'メイリオ')
    p.paragraph_format.space_before=Pt(14); p.paragraph_format.space_after=Pt(6)

def H3(text, color=ORANGE):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = color
    r.font.name='メイリオ'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'メイリオ')
    p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(3)

def P(text, size=10.5, bold=False, color=None, indent=0):
    p = doc.add_paragraph()
    if indent: p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.line_spacing = 1.45; p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); r.font.size = Pt(size); r.font.bold = bold
    if color: r.font.color.rgb = color
    r.font.name='メイリオ'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'メイリオ')

def HR():
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(6)
    r=p.add_run('─'*45); r.font.color.rgb=GRAY; r.font.size=Pt(9)

def shade_table(table, hex_colors):
    """テーブルのセル背景を塗る"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import nsdecls
    from docx.oxml import parse_xml
    for row, color in zip(table.rows, hex_colors):
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
            tcPr.append(shd)

# ============================================================
# 表紙
# ============================================================
H1('石神井南中学校　校内研修')
H1('📦 最終納品物一覧')
P('')
P('日時：令和８年６月３日（水）14:10〜15:30', bold=True, color=NAVY)
P('講師：練馬区教育委員会　指導主事　紺多　章一郎')
P('テーマ：「指導と評価の一体化」〜「主体的に学習に取り組む態度」の評価を中心に〜')
HR()

# サマリー
H2('概要')
P('校長依頼「主体的に学習に取り組む態度の評価のばらつきを防ぐ」を中心に、')
P('講話40分＋研究協議のための教材一式を作成しました。')
P('当日使う中核ファイルは下記4点です。', bold=True, color=ORANGE)

# 一覧表（簡易）
H2('当日使う中核ファイル 4点')
table = doc.add_table(rows=5, cols=3)
table.autofit = False
table.columns[0].width = Cm(0.7)
table.columns[1].width = Cm(8.5)
table.columns[2].width = Cm(6.8)

headers = ['#', 'ファイル名', '用途']
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = h
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
            r.font.name='メイリオ'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'メイリオ')

rows_data = [
    ('1', '講義スライド_..._v13.pptx', '本体スライド（全33枚）'),
    ('2', '講義スライド_..._v13_note更新.pptx', '発表者ビュー用（ノート埋め込み版）'),
    ('3', '発表原稿_v13対応_..._20260603.docx', '紙の発表原稿（口語・約7,300字）'),
    ('4', '質疑応答想定集30問_v13対応_..._20260603.docx', '質疑応答準備（複数視点・約14,900字）'),
]
for i, (num, name, use) in enumerate(rows_data):
    row = table.rows[i+1]
    row.cells[0].text = num
    row.cells[1].text = name
    row.cells[2].text = use
    for c in row.cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10); r.font.name='メイリオ'
                r._element.rPr.rFonts.set(qn('w:eastAsia'),'メイリオ')

# Header row shading
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
header_row = table.rows[0]
for cell in header_row.cells:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="2A96C2"/>')
    tcPr.append(shd)

doc.add_page_break()

# ============================================================
# 各ファイルの詳細
# ============================================================
H1('📄 各ファイルの詳細')

# File 1
H2('① 講義スライド_..._v13.pptx')
H3('用途：プロジェクタ投影用の本体スライド')
P('〔ファイル名〕講義スライド_石神井南中_20260603_v13.pptx', size=10, color=GRAY)
H3('構成（全33枚）', color=DEEP_WATER)
P('• S1-3　表紙／本日の内容／はじめに章扉')
P('• S4　★冒頭クイズ（ペアトーク③）')
P('• S5　数字あてクイズ（区内3校の社会科データ）')
P('• S6　練馬区の評定状況調査結果（評定5のばらつき）')
P('• S7　★「指導と評価の一体化」サイクル')
P('• S8-13　Part1：資質・能力／3つの学び／授業改善と評価は一体')
P('• S14　★ペアトーク①（主体的・対話的・深い学びが見えた場面）')
P('• S15-17　Part2章扉／現状の課題／改善の基本方針')
P('• S18　★授業改善につながる評価（形成的評価サイクル）')
P('• S19-21　3観点の基本構造／知技／思判表')
P('• S22-23　主体的態度の2側面／工夫例')
P('• S24-26　行動観察／発問の工夫／Ｂの姿の数学例')
P('• S27　評価のばらつきを防ぐ（妥当性×信頼性）')
P('• S28　各教科会で揃える3つのポイント（総括例AAA→A、ABB→B等）')
P('• S29　★ペアトーク②（あなたの教科のＢの姿）')
P('• S30　まとめ（明日からできること）')
P('• S31　グループ協議への橋渡し（3つの問い）')
P('• S32-33　今後に向けて：次期論点整理／今日の取り組みは生きる')
H3('特徴', color=DEEP_WATER)
P('・全スライドにメイリオフォント統一')
P('・配色テイスト（水色・紺・オレンジ）で統一')
P('・各スライドに公的資料の出典バー（指導要領／国研／東京都／中教審）')
P('・ペアトーク3枚は緑系の章バーで識別')

# File 2
H2('② 講義スライド_..._v13_note更新.pptx')
H3('用途：発表者ビュー（リハーサル・本番でリアルタイム原稿表示）')
P('〔ファイル名〕講義スライド_石神井南中_20260603_v13_note更新.pptx', size=10, color=GRAY)
H3('内容', color=DEEP_WATER)
P('・①のスライドと内容は同一（見た目は変わらない）')
P('・各スライドに「発表者ノート」として原稿を埋め込み済み')
P('・PowerPointの発表者ビュー（Alt+F5）で起動すると、')
P('　スライドの隣に話す原稿が表示されます')
H3('使い方', color=DEEP_WATER)
P('1. PowerPointで開く')
P('2. リボン「スライドショー」→「発表者ツール」にチェック')
P('3. F5またはAlt+F5で起動')
P('4. 左に大きなスライド／右にノート＋次スライドプレビュー＋経過時間')

# File 3
H2('③ 発表原稿_v13対応_..._20260603.docx')
H3('用途：紙に印刷して手元に置く発表原稿')
P('〔ファイル名〕発表原稿_v13対応_石神井南中_20260603.docx', size=10, color=GRAY)
H3('内容', color=DEEP_WATER)
P('・全33スライドの「タイトル」「目安時間」「話す原稿」を時系列で記載')
P('・口語体（実際にしゃべる言葉）')
P('・約7,300字（A4で約8-10ページ）')
P('・ペアトーク3回（S4／S14／S29）は★マークで識別')
H3('時間配分の目安', color=DEEP_WATER)
P('• 開会〜本題前(S1-3)　約2分')
P('• ★冒頭クイズ(S4)　2分')
P('• 数字・データ(S5-6)　3分')
P('• Part1+ペアトーク①(S7-14)　約13分')
P('• Part2(S15-28)　約24分')
P('• ★ペアトーク②(S29)　2分')
P('• まとめ・今後(S30-33)　約6分')
P('• 合計：約52分（講話40分に収めるには各原稿を適宜短縮）', bold=True, color=ORANGE)

# File 4
H2('④ 質疑応答想定集30問_v13対応_..._20260603.docx')
H3('用途：研修中の質疑応答／グループ協議巡回時の参考')
P('〔ファイル名〕質疑応答想定集30問_v13対応_石神井南中_20260603.docx', size=10, color=GRAY)
H3('構成（全30問・約14,900字）', color=DEEP_WATER)
P('〈第Ⅰ部 講義内容への質疑 15問〉', bold=True, color=NAVY)
P('  Q1-5　サイクル運用／教科会時間／2側面／極端ABC組合せ／妥当性vs信頼性', size=10)
P('  Q6-10　形成的評価と年間計画／総括ルール／次期論点整理／指導と評価の境目／ペアトーク効果', size=10)
P('  Q11-15　観点間の重複／地域差／異動後の継続性／正当な評価／参考資料の活用', size=10)
P('')
P('〈第Ⅱ部 日常の評価で困っていること 15問〉', bold=True, color=NAVY)
P('  Q1-5　不登校／満点生徒の態度／書字困難／外国にルーツ／実技教科', size=10)
P('  Q6-10　保護者クレーム／3と4の境界／ABC換算／特別支援／生成AI', size=10)
P('  Q11-15　TT・少人数／転入生／作問／振り返りの形骸化／総合・特活', size=10)
H3('回答の3段構成', color=DEEP_WATER)
P('① まず端的に　── 実用最優先の答え')
P('② 複数の専門家視点から　── 学習評価研究者／教育心理学者／指導主事／')
P('　 管理職／ベテラン教員／特別支援教育専門家／教科教育専門家 から3-4名分')
P('③ 指導主事としての結論　── 現場で使えるまとめ')

doc.add_page_break()

# ============================================================
# 校長依頼への対応サマリー
# ============================================================
H1('🎯 校長依頼への対応サマリー')
P('')
table2 = doc.add_table(rows=6, cols=2)
table2.columns[0].width = Cm(6.0)
table2.columns[1].width = Cm(10.0)
hdr2 = table2.rows[0]
hdr2.cells[0].text = '校長依頼キーワード'
hdr2.cells[1].text = '対応スライド'
for c in hdr2.cells:
    tcPr = c._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="2A96C2"/>')
    tcPr.append(shd)
    for p in c.paragraphs:
        for r in p.runs:
            r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
            r.font.name='メイリオ'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'メイリオ')

rows_corr = [
    ('指導と評価の一体化', 'S7 サイクル図／S13 一体化は理論'),
    ('正当な評価', 'S27 妥当性×信頼性'),
    ('授業改善につながる評価', 'S18 形成的評価サイクル'),
    ('ばらつきが出ない工夫', 'S6 練馬区データ／S28 教科会で揃える3つ'),
    ('講話40分＋協議への接続', 'S31 グループ協議への橋渡し（3つの問い）'),
]
for i, (k, v) in enumerate(rows_corr):
    row = table2.rows[i+1]
    row.cells[0].text = k
    row.cells[1].text = v
    for c in row.cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10); r.font.name='メイリオ'
                r._element.rPr.rFonts.set(qn('w:eastAsia'),'メイリオ')

# ============================================================
# 当日チェックリスト
# ============================================================
H1('⏰ 当日チェックリスト')
P('')
H3('前日まで', color=ORANGE)
P('☐ PowerPoint「発表者ビュー」の動作確認（Alt+F5）')
P('☐ 発表原稿Wordを印刷して手元に')
P('☐ 質疑応答想定集を手元に（ファイル名がたまにヒットしやすい体裁で）')
P('☐ 練馬区評定状況データの最新版を手元に確認')

H3('当日（14:00 学校到着）', color=ORANGE)
P('☐ プロジェクタ・スピーカー接続確認')
P('☐ スマホ等で「タイマー」アプリ準備（ペアトーク3回分）')
P('☐ 「発表者ビュー」起動テスト')
P('☐ 木原校長への挨拶／田口指導主事（同行）との打合せ')

H3('講話中', color=ORANGE)
P('☐ S4冒頭クイズ：1分相談→挙手で確認→「今日のテーマ」へ伏線')
P('☐ S14ペアトーク①：巡回して具体例を1-2件メモ→還元')
P('☐ S29ペアトーク②：巡回して書けた例を1-2件メモ→協議へ接続')
P('☐ S31グループ協議への橋渡し：3つの問いを明示')

H3('講話後（協議会・質疑応答）', color=ORANGE)
P('☐ 教科会巡回：質疑応答想定集を参照')
P('☐ 質疑応答時：複数視点で答える（指導主事の見解として）')
P('☐ 最後のまとめ：本日のキーメッセージを再強調')

# ============================================================
# GitHub情報
# ============================================================
H1('📁 ファイル取得先（GitHub）')
P('')
P('リポジトリ：contynet730-star/con30', bold=True)
P('ブランチ：claude/stoic-hawking-397Jg', bold=True)
P('')
P('ブラウザで以下を開いてください：', color=NAVY)
P('https://github.com/contynet730-star/con30/tree/claude/stoic-hawking-397Jg',
  size=11, color=DEEP_WATER, bold=True)
P('')
P('各ファイル名をクリックすると、内容プレビューと「Download raw file」ボタンが表示されます。',
  size=10, color=GRAY)

HR()
P('以上。研修当日のご成功をお祈りします。', color=GRAY, bold=True)
P('ご質問・修正等あればいつでもどうぞ。 ─ 紺多 章一郎', color=GRAY)

import os
os.makedirs('/tmp/output', exist_ok=True)
out = '/tmp/output/最終納品物一覧_石神井南中_20260603.docx'
doc.save(out)
print('saved:', out)
print('approx chars:', sum(len(p.text) for p in doc.paragraphs))
