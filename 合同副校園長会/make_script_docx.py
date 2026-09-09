# -*- coding: utf-8 -*-
"""令和8年度9月合同副校園長会（教務関係）発表原稿（推敲版）をWord形式で出力する。"""
import zipfile
from xml.dom import minidom

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DECK = "令和8年度9月合同副校園長会_教務関係.pptx"
OUT = "令和8年度9月合同副校園長会_教務関係_発表原稿（推敲版）.docx"

SLIDE_TITLES = {
    1: "表紙　令和8年度9月合同副校園長会　教務関係について",
    2: "1　実施授業時数について（1学期）",
    3: "2　全国学力・学習状況調査の結果および活用について　－教科に関する調査－",
    4: "2　全国学力・学習状況調査の結果および活用について　－質問紙調査①－",
    5: "2　全国学力・学習状況調査の結果および活用について　－質問紙調査②－",
    6: "3　令和8年度授業改善推進プランの作成および授業改善推進プランに基づく教育活動の実施について",
    7: "4　その他　－東京都中学校英語スピーキングテスト（ESAT-J）－",
}


def notes(deck):
    z = zipfile.ZipFile(deck)
    result = {}
    for i in range(1, 8):
        d = minidom.parseString(z.read(f"ppt/notesSlides/notesSlide{i}.xml"))
        paras = []
        for p in d.getElementsByTagName("a:p"):
            s = "".join(t.firstChild.nodeValue if t.firstChild else "" for t in p.getElementsByTagName("a:t"))
            if s.strip() and s.strip() != str(i):
                paras.append(s)
        result[i] = paras
    return result


def set_font(run, size, bold=False, color=None, name="MS Gothic"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_para(doc, text, size=11, bold=False, color=None, before=0, after=6,
             line=None, indent=None, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = Pt(line)
    if indent:
        pf.left_indent = Cm(indent)
    if text:
        set_font(p.add_run(text), size, bold, color)
    return p


doc = Document()
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21.0)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

add_para(doc, "令和8年度9月　合同副校園長会", 12, bold=True, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "教務関係　発表原稿", 18, bold=True, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "令和8年9月10日（木）　教育指導課　指導主事", 10.5, after=2,
         align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "読み上げ目安　約8分（300字／分）", 9, color=(0x66, 0x66, 0x66), after=14,
         align=WD_ALIGN_PARAGRAPH.CENTER)

data = notes(DECK)
for i in range(1, 8):
    add_para(doc, f"【スライド{i}】{SLIDE_TITLES[i]}", 11, bold=True,
             color=(0x1F, 0x3B, 0x63), before=10 if i > 1 else 0, after=6)
    for para in data[i]:
        add_para(doc, para, 11, after=8, line=20, indent=0.4)

doc.save(OUT)
print("wrote", OUT)
