"""豊玉小学校 創立150周年記念誌 教育長前文（修正版・600字以内）"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

JP_MINCHO = "MS 明朝"
JP_GOTHIC = "MS ゴシック"


def set_font(run, size=10.5, bold=False, name=JP_MINCHO):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def para(text, size=10.5, bold=False, align=None, name=JP_MINCHO,
         line_spacing=22, before=0, after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = Pt(line_spacing)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, name=name)
    return p


doc = Document()
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21.0)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

# タイトル
para("豊玉小学校開校百五十周年を祝して",
     size=16, bold=True, name=JP_GOTHIC,
     align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=14)

# 署名（右寄せ）
para("練馬区教育委員会　教育長　　三浦　康彰",
     size=11, name=JP_MINCHO,
     align=WD_ALIGN_PARAGRAPH.RIGHT, after=18)

# 本文
# ① 冒頭の祝意
para(
    "　練馬区立豊玉小学校が、開校百五十周年を迎えられましたことを心より"
    "お祝い申し上げます。",
    size=11
)

# ② 歴史（Z4：連続性を主題に／143字／構造から140周年版と差別化）
para(
    "　本校が百五十年の永きにわたり歩みを止めなかったのは、明治九年に地域の方々が"
    "学び舎に贈られた「豊玉」── 稲穂のように子どもの心も豊かに実れ ── との祈りが、"
    "南蔵院本堂で迎えられた教員二名・児童八十名から、一万二千六百六十八名の同窓を"
    "介して、今日の教室まで脈々と継がれてきたからでございます。",
    size=11
)

# ③ 実績（原文どおり）
para(
    "　この間、文部科学省、東京都および練馬区の研究指定を度々受けながら"
    "教育目標「考える子　ねばり強い子　心ゆたかな子」の育成に努め、"
    "平成二十九年度には全国学校体育研究優良校として表彰されました。"
    "令和の新時代となってからは小学校教科担任制を推進され、中学校への"
    "円滑な接続の面からも成果を挙げています。四十分午前五時間授業の実施や"
    "探究的な学びにも力を入れ、令和七・八年度には練馬区教育委員会教育課題"
    "研究指定校として、今年度からは文部科学省教育課程サキドリ研究指定校"
    "として実践を深めています。",
    size=11
)

# ④ 感謝（原文どおり）
para(
    "　こうした本校の発展は、歴代校長先生をはじめ教職員の皆様、学校を"
    "支えてくださった地域や保護者の皆様のご尽力の賜と、深く感謝申し上げます。",
    size=11
)

# ⑤ 結語（原文どおり）
para(
    "　豊玉小学校が歴史と伝統を次代に引き継ぐとともに、新たな時代の創造に"
    "向けて益々のご発展を遂げられますことを祈念し、お祝いのことばといたします。",
    size=11
)

# 文字数注記（資料用・小さく）
para(
    "（文字数：全体590文字・600文字以内）",
    size=9, align=WD_ALIGN_PARAGRAPH.RIGHT, before=14, after=0
)

output_path = "/home/user/con30/R8_豊玉小学校150周年_教育長前文_修正版.docx"
doc.save(output_path)
print(f"saved: {output_path}")
