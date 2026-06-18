# -*- coding: utf-8 -*-
"""
今井様あて「学校公開日の日程について（回答）」最新版（プレーン文体・横書き）。

上司レビュー版（プリント）の文体（番号見出しなしの散文）に合わせて再構成し、
以下の内容チェック・修正を反映：
  ② オレンジ部分の言い換え（誰が読んでも分かる平易な表現に）
     旧：「多くの保護者や地域の皆様の利便性を考慮した実施の趣旨につきまして、」
     新：「より多くの保護者や地域の皆様が参加しやすいよう、あらかじめ定めた日に
          実施していることにつきまして、」
  A  文法修正：「実施日は…第二土曜日を基本とすることで…ためです」の主述不整合を
     2文に分割。
  C  「①…こと②…」→「①…こと、②…」読点を補う。
  E  文末に【担当】（連絡先）を追記。
出典：添付3資料／練馬区立学校の管理運営に関する規則
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# ── 本文（散文・段落単位）──────────────────────────────
PARAS = [
    "　お問い合わせいただいた件について、教育指導課より回答いたします。",
    "　土曜日は、区立学校の管理運営に関する規則上、本来「休業日」とされておりますが、"
    "本区では、①必要な授業時数を確保すること、②保護者や地域の皆様に教育活動を広く公開し、"
    "学校教育への理解と信頼を高めること、を意義として、平成24年度から、教育委員会の関与の"
    "下、振替休業日を設定しない土曜授業を全小・中学校で実施しております。",
    "　実施日は、平成24年度以降、原則として第二土曜日を基本としております。これは、"
    "あらかじめ定まった日とすることで、保護者の皆様が参観のご予定を立てやすく、また地域や"
    "区の行事計画の見通しにもなり、より多くの方にご参加いただけるためです。なお、教員の"
    "働き方改革等により、令和6年度からは従来の年8回から、実施月を問わず年間4回（第二土曜日）"
    "に縮減するなどの見直しを行ってきたところです。",
    "　今井様からお申出がありました日程の固定によるご負担につきましては、教職員の多様な"
    "働き方や保護者としての参観機会の確保に関わる重要な視点として受け止めておりますが、"
    "より多くの保護者や地域の皆様が参加しやすいよう、あらかじめ定めた日に実施していること"
    "につきまして、何卒ご理解を賜りますようお願い申し上げます。",
    "　引き続き、練馬区立学校の教育活動にご理解とご協力を賜りますよう、お願い申し上げます。",
]

# 文字数チェック（本文のみ・字下げ除く）
body = "".join(PARAS)
print(f"本文文字数（空白除く）：{len(body.replace('　',''))}字")

# ── 文書生成 ─────────────────────────────────────────
doc = Document()
sec = doc.sections[0]
sec.page_height = Cm(29.7)
sec.page_width = Cm(21.0)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)

MINCHO = "ＭＳ 明朝"
GOTHIC = "ＭＳ ゴシック"


def set_font(run, size, bold=False, name=MINCHO):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def para(text="", size=10.5, bold=False, align=None, name=MINCHO,
         before=0, after=6, line=20):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = Pt(line)
    if align is not None:
        p.alignment = align
    if text:
        set_font(p.add_run(text), size, bold=bold, name=name)
    return p


# 発信日・発信者・宛名
para("令和８年６月　　日", align=WD_ALIGN_PARAGRAPH.RIGHT, after=2)
para("練馬区教育委員会　教育振興部　教育指導課", align=WD_ALIGN_PARAGRAPH.RIGHT, after=12)
para("今井　美咲　様", size=11, after=12)

# 件名
para("学校公開日の日程について（回答）", size=13, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, name=GOTHIC, after=12)

# 本文
for t in PARAS:
    para(t, after=6, line=20)

# 担当
para("", after=4)
para("【担当】", size=10, name=GOTHIC, after=0, line=15)
para("練馬区教育委員会　教育振興部　教育指導課", size=10, after=0, line=15)
para("電話　５９８４－５７５９", size=10, after=0, line=15)

doc.save("/home/user/con30/学校公開日の日程について（回答・改訂版）.docx")
print("Saved: 学校公開日の日程について（回答・改訂版）.docx")
