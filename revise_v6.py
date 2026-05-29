# -*- coding: utf-8 -*-
"""
アップロード版 v5（27枚）をベースに、専門家指摘の7点を反映して30枚に改修。
1 タイトル変更（指導と評価の一体化〜主体的態度を中心に〜）
2 「指導と評価の一体化」サイクル図を冒頭に追加
3 3観点を一体化の中で対等に位置づけ（新フレーム内で明示）
4 「授業改善につながる評価」＝形成的評価のスライド新設
5 目標→指導→評価→改善の単元設計フロー（点2のサイクル内に集約）
6 グループ協議への橋渡しスライド追加
7 S26「9つの働きかけ」参照切れ→実在内容へ修正
＋ ズレたページ番号(S7/S10)除去、フォントをメイリオ統一
"""
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

# ---- palette (v5と一致を確認済み) ----
WATER_BLUE = RGBColor(0x5B, 0xC0, 0xDE)
LIGHT_WATER= RGBColor(0xD6, 0xEE, 0xFA)
DEEP_WATER = RGBColor(0x2A, 0x96, 0xC2)
NAVY       = RGBColor(0x1F, 0x3A, 0x6E)
ORANGE_KW  = RGBColor(0xEB, 0x6C, 0x15)
RED        = RGBColor(0xC0, 0x40, 0x40)
GREEN      = RGBColor(0x07, 0xA9, 0x73)
YELLOW     = RGBColor(0xFF, 0xE0, 0x66)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
BLACK      = RGBColor(0x00, 0x00, 0x00)
DARK_GRAY  = RGBColor(0x40, 0x40, 0x40)
F_MAIN = "メイリオ"

prs = Presentation('/tmp/conv/slides.pptx')
SW, SH = prs.slide_width, prs.slide_height

# ============================================================
# helpers
# ============================================================
def blank():
    return prs.slides.add_slide(prs.slide_layouts[6])

def chapter_bar(slide, text, fontsize=22):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Emu(720000))
    bar.fill.solid(); bar.fill.fore_color.rgb = DEEP_WATER; bar.line.fill.background()
    tf = bar.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(300000); tf.margin_right=Emu(100000); tf.margin_top=Emu(20000); tf.margin_bottom=Emu(20000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = text
    r.font.size=Pt(fontsize); r.font.bold=True; r.font.color.rgb=WHITE; r.font.name=F_MAIN

def sub_heading(slide, text, fontsize=28):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(830000), SW, Emu(700000))
    bar.fill.solid(); bar.fill.fore_color.rgb=LIGHT_WATER; bar.line.fill.background()
    tf=bar.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(300000); tf.margin_right=Emu(200000); tf.margin_top=Emu(20000); tf.margin_bottom=Emu(20000)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT
    r=p.add_run(); r.text=text
    r.font.size=Pt(fontsize); r.font.bold=True; r.font.color.rgb=NAVY; r.font.name=F_MAIN
    ul=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(1530000), SW, Emu(40000))
    ul.fill.solid(); ul.fill.fore_color.rgb=WATER_BLUE; ul.line.fill.background()

def box(slide, left, top, width, height, lines, default_size=24, default_color=BLACK,
        line_spacing=1.3, align=PP_ALIGN.LEFT, fill_color=WHITE, border_color=WATER_BLUE,
        border_width=3.0, vanchor=MSO_ANCHOR.MIDDLE, default_bold=True, shape=MSO_SHAPE.RECTANGLE):
    b = slide.shapes.add_shape(shape, left, top, width, height)
    b.fill.solid(); b.fill.fore_color.rgb=fill_color
    if border_color: b.line.color.rgb=border_color; b.line.width=Pt(border_width)
    else: b.line.fill.background()
    tf=b.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(140000); tf.margin_right=Emu(140000); tf.margin_top=Emu(80000); tf.margin_bottom=Emu(80000)
    tf.vertical_anchor=vanchor
    for i,item in enumerate(lines):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if isinstance(item, list):
            for j,ri in enumerate(item):
                rt,ro = ri if isinstance(ri,tuple) else (ri,{})
                if j==0:
                    p.alignment=ro.get('align',align); p.line_spacing=ro.get('line_spacing',line_spacing)
                r=p.add_run(); r.text=rt
                r.font.size=Pt(ro.get('size',default_size)); r.font.bold=ro.get('bold',default_bold)
                r.font.color.rgb=ro.get('color',default_color); r.font.name=F_MAIN
        else:
            text,o = item if isinstance(item,tuple) else (item,{})
            p.alignment=o.get('align',align); p.line_spacing=o.get('line_spacing',line_spacing)
            r=p.add_run(); r.text=text
            r.font.size=Pt(o.get('size',default_size)); r.font.bold=o.get('bold',default_bold)
            r.font.color.rgb=o.get('color',default_color); r.font.name=F_MAIN
    return b

def arrow_right(slide, left, top, width, height, color=ORANGE_KW):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    a.fill.solid(); a.fill.fore_color.rgb=color; a.line.fill.background()
    return a

# 4-box equal layout x-coords (used by new1/new2)
BOX_W = Emu(2550000); GAP = Emu(400000)
XS = [Emu(400000), Emu(3350000), Emu(6300000), Emu(9250000)]
ARR_X = [Emu(2950000), Emu(5900000), Emu(8850000)]

# ============================================================
# (1) S1 タイトル変更
# ============================================================
s1 = prs.slides[0]
title = s1.shapes[0].text_frame.paragraphs[0]
title.runs[0].text = "「指導と評価の一体化」"
for r in title.runs[1:]:
    r.text = ""
# サブタイトル：主体的態度を「中心に」へ
s1.shapes[1].text_frame.paragraphs[0].runs[0].text = "〜「主体的に学習に取り組む態度」の評価を中心に〜"

# ============================================================
# (7) S26 「9つの働きかけ」参照切れ → 実在する工夫へ
# ============================================================
s26 = prs.slides[25]
for sh in s26.shapes:
    if sh.has_text_frame and "9つの働きかけ" in sh.text_frame.text:
        sh.text_frame.paragraphs[0].runs[0].text = "□ 評価の工夫（行動観察・発問・自己評価）から1つ選んで授業に入れる"

# ============================================================
# ズレたページ番号(S7"8"/S10"11")を除去
# ============================================================
for s in prs.slides:
    for sh in list(s.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit():
            if sh.left and sh.left > Emu(10000000) and sh.top and sh.top > Emu(6000000):
                sh._element.getparent().remove(sh._element)

# ============================================================
# NEW1: 「指導と評価の一体化」とは（サイクル：点2+5+3）
# ============================================================
n1 = blank()
chapter_bar(n1, "０　はじめに")
sub_heading(n1, "「指導と評価の一体化」とは ─ 本研修の土台")
cycle = [
    ("① 目標", "「Ｂの姿」を描く"),
    ("② 指導", "授業を行う\n（主体的・対話的で深い学び）"),
    ("③ 評価", "学びの姿をみとる"),
    ("④ 改善", "次の指導に生かす"),
]
for i,(head,desc) in enumerate(cycle):
    box(n1, XS[i], Emu(2000000), BOX_W, Emu(1450000), [
        (head, {'size':24,'bold':True,'color':WHITE,'align':PP_ALIGN.CENTER}),
        (desc, {'size':16,'color':NAVY,'align':PP_ALIGN.CENTER}),
    ], fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0, line_spacing=1.15)
for ax in ARR_X:
    arrow_right(n1, ax, Emu(2520000), Emu(400000), Emu(420000))
box(n1, Emu(400000), Emu(3650000), Emu(11400000), Emu(620000), [
    [("④の評価を ①②③ に生かして単元の中でくり返す", {'size':20,'color':NAVY}),
     ("　＝　これが「指導と評価の一体化」", {'size':22,'color':ORANGE_KW})],
], fill_color=LIGHT_WATER, border_width=3.0, align=PP_ALIGN.CENTER)
box(n1, Emu(400000), Emu(4420000), Emu(11400000), Emu(1450000), [
    [("「正当な評価」", {'size':22,'color':ORANGE_KW}),
     ("のために授業と評価を一体で設計し、", {'size':20,'color':BLACK}),
     ("「授業改善につながる評価」", {'size':22,'color':ORANGE_KW}),
     ("にする", {'size':20,'color':BLACK})],
    [("本研修は", {'size':18,'color':BLACK}),
     ("3観点すべて", {'size':20,'color':NAVY}),
     ("を対象とし、特に見取りが難しい", {'size':18,'color':BLACK}),
     ("「主体的に学習に取り組む態度」を重点", {'size':20,'color':NAVY}),
     ("的に扱う", {'size':18,'color':BLACK})],
], border_width=3.5, align=PP_ALIGN.CENTER, line_spacing=1.35)

# ============================================================
# NEW2: 授業改善につながる評価（形成的評価：点4）
# ============================================================
n2 = blank()
chapter_bar(n2, "２　『主体的に学習に取り組む態度』の学習評価")
sub_heading(n2, "授業改善につながる評価 ─ 評価を次の指導に生かす")
box(n2, Emu(400000), Emu(1850000), Emu(11400000), Emu(720000), [
    [("評価には2つの役割：", {'size':20,'color':BLACK}),
     ("①記録・評定のため（総括的）", {'size':19,'color':DARK_GRAY}),
     ("　／　", {'size':19,'color':BLACK}),
     ("②指導改善のため（形成的）", {'size':20,'color':ORANGE_KW}),
     ("　← 校長が重視", {'size':19,'color':RED})],
], fill_color=LIGHT_WATER, border_width=3.0, align=PP_ALIGN.CENTER)
form = [
    ("① みとる", "評価で学習状況を把握"),
    ("② 気づく", "つまずき・伸びを発見"),
    ("③ 変える", "次時の指導を修正"),
    ("④ 伸びる", "生徒の学習が改善"),
]
for i,(head,desc) in enumerate(form):
    box(n2, XS[i], Emu(2750000), BOX_W, Emu(1250000), [
        (head, {'size':23,'bold':True,'color':WHITE,'align':PP_ALIGN.CENTER}),
        (desc, {'size':15,'color':NAVY,'align':PP_ALIGN.CENTER}),
    ], fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0, line_spacing=1.15)
for ax in ARR_X:
    arrow_right(n2, ax, Emu(3230000), Emu(400000), Emu(360000))
box(n2, Emu(400000), Emu(4180000), Emu(11400000), Emu(1080000), [
    ("＜例＞", {'size':18,'color':ORANGE_KW,'bold':True}),
    ("振り返りで「見通しが弱い」生徒が多い → 次時の冒頭に「今日のゴールと手順」を確認する場面を加える",
     {'size':19,'color':NAVY}),
], border_width=3.0, line_spacing=1.3, align=PP_ALIGN.LEFT)
box(n2, Emu(400000), Emu(5380000), Emu(11400000), Emu(560000), [
    [("評価は", {'size':21,'color':BLACK}),
     ("“つける”だけでなく“生かす”", {'size':23,'color':ORANGE_KW}),
     ("。これが校長の言う「授業改善につながる評価」", {'size':21,'color':BLACK})],
], fill_color=YELLOW, border_color=ORANGE_KW, border_width=3.0, align=PP_ALIGN.CENTER)

# ============================================================
# NEW3: グループ協議への橋渡し（点6）
# ============================================================
n3 = blank()
chapter_bar(n3, "グループ協議へ")
sub_heading(n3, "このあと、各教科の評価資料を持ち寄って検討します")
box(n3, Emu(400000), Emu(1820000), Emu(11400000), Emu(720000), [
    [("協議テーマ：", {'size':20,'color':BLACK}),
     ("各教科で評価方法の「ばらつき」を減らす", {'size':22,'color':ORANGE_KW})],
], fill_color=LIGHT_WATER, border_width=3.0, align=PP_ALIGN.CENTER)
qs = [
    ("①", "あなたの教科で「Ｂと判断する姿」は、具体的にどんな姿か？"),
    ("②", "単元のどの場面で、どんな方法でみとるか？"),
    ("③", "その評価を、次の授業改善にどう生かすか？"),
]
y = 2720000
for i,(num,q) in enumerate(qs):
    yy = Emu(y + i*900000)
    box(n3, Emu(400000), yy, Emu(1000000), Emu(780000), [
        (num, {'size':36,'bold':True,'color':WHITE,'align':PP_ALIGN.CENTER}),
    ], fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
    box(n3, Emu(1500000), yy, Emu(10300000), Emu(780000), [
        (q, {'size':21,'color':NAVY,'align':PP_ALIGN.LEFT}),
    ], border_width=3.0)
box(n3, Emu(400000), Emu(5520000), Emu(11400000), Emu(620000), [
    ("持ち寄った評価資料をもとに話し合い、教科として評価方法をそろえる",
     {'size':21,'bold':True,'color':NAVY,'align':PP_ALIGN.CENTER}),
], fill_color=YELLOW, border_color=ORANGE_KW, border_width=3.0)

# ============================================================
# スライド並べ替え（27+3=30）
# ============================================================
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)  # 0..29 (27 orig + n1=27,n2=28,n3=29)
desired = [0,1,2,3, 27, 4,5,6,7,8,9, 10,11,12, 28, 13,14,15,16,17,18,19,20, 21,22,23,24, 29, 25,26]
assert sorted(desired)==list(range(30)), "順序不正"
for el in ids: sldIdLst.remove(el)
for i in desired: sldIdLst.append(ids[i])

# ============================================================
# フォントをメイリオに統一（グループ内も再帰）
# ============================================================
from lxml import etree
NS_A='http://schemas.openxmlformats.org/drawingml/2006/main'
def force_meiryo(run):
    rPr = run._r.find(f'{{{NS_A}}}rPr')
    if rPr is None:
        rPr = etree.SubElement(run._r, f'{{{NS_A}}}rPr'); run._r.insert(0, rPr)
    for tag in ['latin','ea','cs']:
        el = rPr.find(f'{{{NS_A}}}{tag}')
        if el is None: el = etree.SubElement(rPr, f'{{{NS_A}}}{tag}')
        el.set('typeface','メイリオ')
def walk_fonts(shapes):
    n=0
    for sh in shapes:
        if sh.shape_type==MSO_SHAPE_TYPE.GROUP:
            n+=walk_fonts(sh.shapes)
        elif sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    force_meiryo(r); n+=1
        elif getattr(sh,'has_table',False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    for p in c.text_frame.paragraphs:
                        for r in p.runs:
                            force_meiryo(r); n+=1
    return n
total = sum(walk_fonts(s.shapes) for s in prs.slides)
print("font runs unified:", total)

import os
os.makedirs('/tmp/output', exist_ok=True)
OUT='/tmp/output/講義スライド_石神井南中_20260603_v6.pptx'
prs.save(OUT)
print("saved:", OUT, "slides:", len(prs.slides))
