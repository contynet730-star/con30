# -*- coding: utf-8 -*-
"""
v7(27枚) → v8(28枚)：3つのPDF資料を根拠として織り込む（プラスになる構想）
1 S5 一体化サイクル：出典バー追加（中学校学習指導要領 解説 総則編 第3章第2節2）
2 S15 授業改善につながる評価：出典バー追加（指導要領＋Q&A実践書 Q5）
3 S25 教科会で揃える3つ：③に「総括の例」ミニバー追加（Q9表より）
4 末尾(S28)に参考文献スライド復活
※ テーマ拡散を避け、新章は増やさない。既存スライドの根拠を強化する形に留める。
"""
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

WATER_BLUE=RGBColor(0x5B,0xC0,0xDE); LIGHT_WATER=RGBColor(0xD6,0xEE,0xFA)
DEEP_WATER=RGBColor(0x2A,0x96,0xC2); NAVY=RGBColor(0x1F,0x3A,0x6E)
ORANGE_KW=RGBColor(0xEB,0x6C,0x15); YELLOW=RGBColor(0xFF,0xE0,0x66)
WHITE=RGBColor(0xFF,0xFF,0xFF); BLACK=RGBColor(0x00,0x00,0x00); DARK_GRAY=RGBColor(0x40,0x40,0x40)
F_MAIN="メイリオ"

prs = Presentation('/tmp/output/講義スライド_石神井南中_20260603_v7.pptx')
SW,SH = prs.slide_width, prs.slide_height

def add_source_line(slide, text, top=Emu(5990000), height=Emu(330000)):
    tb = slide.shapes.add_textbox(Emu(300000), top, Emu(11400000), height)
    tf = tb.text_frame
    tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(40000); tf.margin_right=Emu(40000); tf.margin_top=Emu(10000); tf.margin_bottom=Emu(10000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = "出典：" + text
    r.font.size = Pt(13); r.font.color.rgb = DARK_GRAY; r.font.name = F_MAIN

def add_chapter_bar(slide, text, fs=22):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Emu(720000))
    bar.fill.solid(); bar.fill.fore_color.rgb=DEEP_WATER; bar.line.fill.background()
    tf=bar.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(300000); tf.margin_right=Emu(100000); tf.margin_top=Emu(20000); tf.margin_bottom=Emu(20000)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT
    r=p.add_run(); r.text=text
    r.font.size=Pt(fs); r.font.bold=True; r.font.color.rgb=WHITE; r.font.name=F_MAIN

def add_sub_heading(slide, text, fs=28):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(830000), SW, Emu(700000))
    bar.fill.solid(); bar.fill.fore_color.rgb=LIGHT_WATER; bar.line.fill.background()
    tf=bar.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(300000); tf.margin_right=Emu(200000); tf.margin_top=Emu(20000); tf.margin_bottom=Emu(20000)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT
    r=p.add_run(); r.text=text
    r.font.size=Pt(fs); r.font.bold=True; r.font.color.rgb=NAVY; r.font.name=F_MAIN
    ul=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(1530000), SW, Emu(40000))
    ul.fill.solid(); ul.fill.fore_color.rgb=WATER_BLUE; ul.line.fill.background()

def add_box(slide,left,top,width,height,lines,default_size=24,default_color=BLACK,
            line_spacing=1.3,align=PP_ALIGN.LEFT,fill_color=WHITE,border_color=WATER_BLUE,
            border_width=3.0,vanchor=MSO_ANCHOR.MIDDLE,default_bold=True):
    b=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,left,top,width,height)
    b.fill.solid(); b.fill.fore_color.rgb=fill_color
    if border_color: b.line.color.rgb=border_color; b.line.width=Pt(border_width)
    else: b.line.fill.background()
    tf=b.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(140000); tf.margin_right=Emu(140000); tf.margin_top=Emu(60000); tf.margin_bottom=Emu(60000)
    tf.vertical_anchor=vanchor
    for i,item in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if isinstance(item,list):
            for j,ri in enumerate(item):
                rt,ro = ri if isinstance(ri,tuple) else (ri,{})
                if j==0: p.alignment=ro.get('align',align); p.line_spacing=ro.get('line_spacing',line_spacing)
                r=p.add_run(); r.text=rt
                r.font.size=Pt(ro.get('size',default_size)); r.font.bold=ro.get('bold',default_bold)
                r.font.color.rgb=ro.get('color',default_color); r.font.name=F_MAIN
        else:
            text,o = item if isinstance(item,tuple) else (item,{})
            p.alignment=o.get('align',align); p.line_spacing=o.get('line_spacing',line_spacing)
            r=p.add_run(); r.text=text
            r.font.size=Pt(o.get('size',default_size)); r.font.bold=o.get('bold',default_bold)
            r.font.color.rgb=o.get('color',default_color); r.font.name=F_MAIN
    return b

# ============================================================
# (1) S5 一体化サイクル：S5は最下部空きありbot=5870000→出典を5990000に
# ============================================================
add_source_line(prs.slides[4], "中学校学習指導要領 解説 総則編 第3章第2節2「学習評価の充実」／文部科学省")

# ============================================================
# (2) S15 授業改善につながる評価：bot=5940000 → ぎりぎりなので6080000に
# ============================================================
add_source_line(prs.slides[14],
    "中学校学習指導要領 解説 総則編 第3章第2節2(1)指導の評価と改善／学習評価Q&A実践書 Q5「評価を事務作業にしないために」",
    top=Emu(6080000), height=Emu(310000))

# ============================================================
# (3) S25 教科会で揃える3つ：③カード（bot=5600000）の下に「総括の例」ミニバー追加
# ============================================================
s25 = prs.slides[24]
# 既存の出典バー（T=6430000）を一旦削除して、ミニ例ボックス→出典の順に再配置
for sh in list(s25.shapes):
    if sh.has_text_frame and sh.text_frame.text.startswith("出典："):
        sh._element.getparent().remove(sh._element)

# ミニ例ボックス（T=5650000, H=520000）
add_box(s25, Emu(1400000), Emu(5650000), Emu(10400000), Emu(520000),
    [[("③の例：", {'size':16,'color':DARK_GRAY,'bold':True}),
      ("AAA→A", {'size':17,'color':ORANGE_KW}),
      ("　／　ABB→B", {'size':17,'color':NAVY}),
      ("　／　BBC→B", {'size':17,'color':NAVY}),
      ("　／　CCC→C", {'size':17,'color':NAVY}),
      ("　（教科会で総括ルールを揃える）", {'size':14,'color':DARK_GRAY})]],
    fill_color=LIGHT_WATER, border_width=2.0, line_spacing=1.1, align=PP_ALIGN.LEFT)

# 出典を再配置（複数出典）
add_source_line(s25,
    "東京都教育委員会「指導と評価の一体化を目指して」／学習評価Q&A実践書 Q9「観点別評価から評定への総括」",
    top=Emu(6230000), height=Emu(540000))

# ============================================================
# (4) 末尾に参考文献スライドを追加（S28）
# ============================================================
n_ref = prs.slides.add_slide(prs.slide_layouts[6])
add_chapter_bar(n_ref, "参考文献")
add_sub_heading(n_ref, "本研修の根拠資料")

refs = [
    "● 中学校学習指導要領 解説　総則編　文部科学省",
    "　 第3章第2節2「学習評価の充実」（指導の評価と改善／学習評価に関する工夫）",
    "● 学習評価の在り方ハンドブック（中学校編）",
    "　 文部科学省　国立教育政策研究所教育課程研究センター",
    "● 「指導と評価の一体化」のための学習評価に関する参考資料",
    "　 文部科学省　国立教育政策研究所教育課程研究センター",
    "● 「子供たちに未来の創り手となるために必要な資質・能力を育む",
    "　 指導と評価の一体化を目指して」Ⅰ理論編　Ⅱ実践編　東京都教育委員会",
    "● 学習評価Q&A実践書（観点別評価・パフォーマンス評価ほか）",
    "　 ※形成的評価・総括ルール・パフォーマンス評価の実践解説書",
]
add_box(n_ref, Emu(400000), Emu(1850000), Emu(11400000), Emu(4200000),
        [(r, {'size':20,'color':NAVY,'align':PP_ALIGN.LEFT}) for r in refs],
        line_spacing=1.5, border_width=3.0, fill_color=LIGHT_WATER, default_bold=False,
        vanchor=MSO_ANCHOR.MIDDLE)

# ---- フォント統一 ----
from lxml import etree
NS_A='http://schemas.openxmlformats.org/drawingml/2006/main'
def force(run):
    rPr=run._r.find(f'{{{NS_A}}}rPr')
    if rPr is None: rPr=etree.SubElement(run._r,f'{{{NS_A}}}rPr'); run._r.insert(0,rPr)
    for t in ['latin','ea','cs']:
        e=rPr.find(f'{{{NS_A}}}{t}')
        if e is None: e=etree.SubElement(rPr,f'{{{NS_A}}}{t}')
        e.set('typeface','メイリオ')
def wf(shapes):
    n=0
    for sh in shapes:
        if sh.shape_type==MSO_SHAPE_TYPE.GROUP: n+=wf(sh.shapes)
        elif sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs: force(r); n+=1
        elif getattr(sh,'has_table',False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    for p in c.text_frame.paragraphs:
                        for r in p.runs: force(r); n+=1
    return n
print("font runs:", sum(wf(s.shapes) for s in prs.slides))

import os; os.makedirs('/tmp/output',exist_ok=True)
OUT='/tmp/output/講義スライド_石神井南中_20260603_v8.pptx'
prs.save(OUT); print("saved:",OUT,"slides:",len(prs.slides))
