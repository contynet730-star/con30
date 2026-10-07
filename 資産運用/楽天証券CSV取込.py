"""楽天証券「保有商品一覧」CSVを 資産運用シミュレーション.xlsx の「保有資産」シートに取り込む。

使い方:
    python 楽天証券CSV取込.py 保有商品一覧.csv 資産運用シミュレーション.xlsx

CSVの見出し行（「銘柄」と「時価評価額」を含む行）を探し、列名で値を対応付ける。
列の並びが変わっても、見出しの名前が同じなら取り込める。
"""
import csv
import datetime
import io
import re
import sys

from openpyxl import load_workbook

SHEET = "保有資産"
FIRST, LAST = 11, 210
# 保有資産シートの列 → CSV見出しの判定
COLUMNS = {
    "A": lambda h: "種別" in h,
    "B": lambda h: "コード" in h or "ティッカー" in h,
    "C": lambda h: h in ("銘柄", "銘柄名", "ファンド名"),
    "D": lambda h: "口座" in h,
    "E": lambda h: "保有数量" in h or h == "数量",
    "F": lambda h: "平均取得" in h,
    "G": lambda h: h.startswith("現在値") and "更新" not in h,
    "H": lambda h: "時価評価額" in h,
    "I": lambda h: "評価損益" in h and "%" not in h and "％" not in h,
}
NUMERIC = set("EFGHI")


def read_rows(path):
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "cp932"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    return list(csv.reader(io.StringIO(text)))


def to_number(s):
    s = re.sub(r"[,\s円口株+＋]", "", s or "")
    if s in ("", "-", "－", "--"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse(rows):
    holdings, mapping = [], None
    for row in rows:
        cells = [c.strip() for c in row]
        if any(c in ("銘柄", "銘柄名", "ファンド名") for c in cells) and any("時価評価額" in c for c in cells):
            mapping = {}
            for col, match in COLUMNS.items():
                for i, h in enumerate(cells):
                    if match(h):
                        mapping[col] = i
                        break
            continue
        if not mapping or "H" not in mapping or "C" not in mapping:
            continue
        get = lambda col: cells[mapping[col]] if col in mapping and mapping[col] < len(cells) else ""
        name, value = get("C"), to_number(get("H"))
        if not name or value is None or "合計" in name:
            continue
        holdings.append({col: (to_number(get(col)) if col in NUMERIC else get(col)) for col in COLUMNS})
    return holdings


def main(csv_path, xlsx_path):
    holdings = parse(read_rows(csv_path))
    if not holdings:
        sys.exit("CSVから保有商品を読み取れませんでした。「保有商品一覧」のCSVか確認してください。")
    if len(holdings) > LAST - FIRST + 1:
        sys.exit(f"銘柄数が上限（{LAST - FIRST + 1}件）を超えています。")

    wb = load_workbook(xlsx_path)
    ws = wb[SHEET]
    for r in range(FIRST, LAST + 1):
        for col in COLUMNS:
            ws[f"{col}{r}"].value = None
    for r, h in enumerate(holdings, start=FIRST):
        for col, v in h.items():
            ws[f"{col}{r}"].value = v
    ws["B2"].value = datetime.date.today()
    wb.calculation.fullCalcOnLoad = True
    wb.save(xlsx_path)

    total = sum(h["H"] or 0 for h in holdings)
    print(f"{len(holdings)}銘柄を取り込みました（時価評価額合計 {total:,.0f}円）")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
