# -*- coding: utf-8 -*-
"""校長会説明資料 発表原稿（Word）を生成する"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import script_data as D

JP = "Meiryo"
INK   = (0x14, 0x31, 0x2A)
PINE  = (0x1F, 0x5D, 0x4C)
MOSSD = (0x5C, 0x8A, 0x4A)
GOLD  = (0x8A, 0x54, 0x12)
GREY  = (0x5B, 0x66, 0x60)
BODY  = (0x1C, 0x23, 0x20)

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin = sec.bottom_margin = Cm(1.8)
sec.left_margin = sec.right_margin = Cm(1.9)

st = doc.styles["Normal"]
st.font.name = JP
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), JP)


def run_font(r, size=11, bold=False, color=None, name=JP):
    r.font.name = name
    r.font.size = Pt(size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    return r


def para(text="", size=11, bold=False, color=BODY, before=0, after=6,
         line=None, indent=0, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if line:
        pf.line_spacing = Pt(line)
    if indent:
        pf.left_indent = Cm(indent)
    if align is not None:
        p.alignment = align
    if text:
        run_font(p.add_run(text), size, bold, color)
    return p


def shade(p, hexcolor):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    p._p.get_or_add_pPr().append(shd)


def cell_shade(c, hexcolor):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    c._tc.get_or_add_tcPr().append(shd)


def band(text, fill="1F5D4C", color=(255, 255, 255), size=13, before=14, after=8):
    p = para(before=before, after=after)
    pf = p.paragraph_format
    pf.left_indent = Cm(0.2)
    run_font(p.add_run(text), size, True, color)
    shade(p, fill)
    p.paragraph_format.keep_with_next = True
    return p


def section_title(text):
    p = para(before=16, after=10)
    run_font(p.add_run(text), 16, True, INK)
    return p


# ── 表紙 ────────────────────────────────────────────
p = para(before=30, after=2)
run_font(p.add_run("校長会説明資料"), 12, True, MOSSD)
p = para(after=4)
run_font(p.add_run("校長権限でつくる"), 24, True, INK)
p = para(after=12)
run_font(p.add_run("令和9年度の学校づくり"), 24, True, INK)
p = para(after=20)
run_font(p.add_run("発 表 原 稿"), 15, True, GOLD)

t = doc.add_table(rows=len(D.OVERVIEW), cols=2)
t.autofit = False
for i, (k, v) in enumerate(D.OVERVIEW):
    row = t.rows[i]
    row.cells[0].width, row.cells[1].width = Cm(3.0), Cm(14.2)
    cell_shade(row.cells[0], "E4EDE6")
    for c, (txt, bold, col, sz) in zip(row.cells, [(k, True, PINE, 10.5), (v, False, BODY, 10.5)]):
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        pp = c.paragraphs[0]
        pp.paragraph_format.space_before = Pt(4)
        pp.paragraph_format.space_after = Pt(4)
        run_font(pp.add_run(txt), sz, bold, col)

para(before=18, after=4, text="読むときの要点", size=13, bold=True, color=INK)
for i, (h, b) in enumerate(D.TIPS, 1):
    p = para(after=2, indent=0.3)
    run_font(p.add_run(f"{i}. {h}"), 11.5, True, PINE)
    para(b, size=10.5, color=BODY, after=8, indent=0.8, line=17)

p = para(before=16, after=0, size=9.5, color=GREY,
         text="※ 本原稿は発表者用です。時刻表示は目安で、スライド1を0分として計算しています。"
              "【　】はト書きです。声に出す部分ではありません。")

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ── 本編原稿 ────────────────────────────────────────
section_title("本編原稿（スライド1〜50）")
para("以下、スライドを送りながら読み上げます。【　】はト書きで、声に出す部分ではありません。",
     size=10.5, color=GREY, after=10)

for no, tm, title, paras, stage in D.SCRIPT:
    band(f"スライド {no}　［{tm}］　{title}", size=12, before=13, after=7)
    for tx in paras:
        para(tx, size=11.5, color=BODY, after=7, line=21, indent=0.2)
    if stage:
        p = para(before=2, after=10, indent=0.2)
        run_font(p.add_run("【 " + stage + " 】"), 10, True, GOLD)
        shade(p, "F4F1E8")

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ── 質疑応答 ────────────────────────────────────────
section_title("質疑応答（スライド51〜56）")
for tx in D.QA_INTRO:
    para(tx, size=11, color=BODY, after=6, line=19, indent=0.2)

para(before=10, after=6, size=11.5, bold=True, color=PINE,
     text="はじめに、この一言を置いてから質疑に入ると場が和みます。")
p = para(after=12, indent=0.5)
run_font(p.add_run("「新しいことを始めるときは、必ず反発が出ます。出ないほうがむしろ危険です。"
                   "反対のご意見ほど、設計を鍛えてくれますので、遠慮なくお聞かせください。」"),
         11.5, False, INK)
shade(p, "F1F5F2")

qn_ = 0
for code, name, items in D.QA:
    band(f"{code}．{name}", fill="5C8A4A", size=12, before=13, after=7)
    for q, a in items:
        qn_ += 1
        p = para(after=3, indent=0.2)
        p.paragraph_format.keep_with_next = True
        run_font(p.add_run(f"Q{qn_}　"), 11, True, GREY)
        run_font(p.add_run(q), 11.5, True, INK)
        p = para(after=9, indent=0.9, line=19)
        run_font(p.add_run("A　"), 11, True, PINE)
        run_font(p.add_run(a), 11, False, BODY)

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ── 30分短縮版 ──────────────────────────────────────
section_title("30分に短縮する場合")
para(D.SHORT_VERSION["lead"], size=11, color=BODY, after=10, line=19, indent=0.2)

t = doc.add_table(rows=len(D.SHORT_VERSION["skip"]) + 1, cols=3)
t.autofit = False
hdr = t.rows[0]
for c, txt, w in zip(hdr.cells, ["飛ばすスライド", "内容", "扱い方"], [Cm(3.0), Cm(6.2)]+[Cm(8.0)]):
    c.width = w
    cell_shade(c, "1F5D4C")
    pp = c.paragraphs[0]
    pp.paragraph_format.space_before = Pt(3)
    pp.paragraph_format.space_after = Pt(3)
    run_font(pp.add_run(txt), 10.5, True, (255, 255, 255))
for i, (sl, nm, how) in enumerate(D.SHORT_VERSION["skip"], 1):
    row = t.rows[i]
    if i % 2 == 1:
        for c in row.cells:
            cell_shade(c, "F1F5F2")
    for c, txt, w, bold, col in zip(row.cells, [sl, nm, how],
                                    [Cm(3.0), Cm(6.2), Cm(8.0)],
                                    [True, False, False],
                                    [PINE, INK, BODY]):
        c.width = w
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        pp = c.paragraphs[0]
        pp.paragraph_format.space_before = Pt(3)
        pp.paragraph_format.space_after = Pt(3)
        run_font(pp.add_run(txt), 10, bold, col)

p = para(before=14, after=6, indent=0.2)
run_font(p.add_run(D.SHORT_VERSION["keep"]), 11.5, True, INK)
shade(p, "E4EDE6")

para(before=14, after=4, size=12, bold=True, color=INK, text="出典")
para("中央教育審議会 初等中等教育分科会 教育課程部会 教育課程企画特別部会（第17回）"
     "令和8年8月31日 資料1「次期学習指導要領等に向けた審議まとめ（素案）」（文部科学省）"
     "ならびに同部会の審議経過に関する各種報道・解説記事",
     size=10, color=GREY, after=4, line=17, indent=0.2)

OUT = sys.argv[1] if len(sys.argv) > 1 else "発表原稿.docx"
doc.save(OUT)
print("written:", OUT)
