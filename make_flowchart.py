# -*- coding: utf-8 -*-
"""
児童生徒の自殺企図に関する相談対応フロー図（こたエール経由）を PowerPoint で生成するスクリプト。
A4縦1枚・縦型フロー（上から下）・2列構成。
  左列: フェーズ1 覚知〜児童生徒の特定
  右列: フェーズ2〜3 警察との連携〜安否確認・事後の支援
"""
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE as MSO_LINE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

JP_FONT = "Meiryo UI"
TEXT_DARK = RGBColor(0x1F, 0x29, 0x37)
LINE_DARK = RGBColor(0x37, 0x47, 0x5A)
BADGE_FILL = RGBColor(0x2F, 0x54, 0x96)
NOTE_EDGE = RGBColor(0x8A, 0x93, 0x9E)
NOTE_FILL = RGBColor(0xF4, 0xF6, 0xF8)
WARN_EDGE = RGBColor(0xB0, 0x21, 0x21)
WARN_FILL = RGBColor(0xFD, 0xEC, 0xEC)
ORANGE = RGBColor(0xC0, 0x63, 0x0C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# 主体ごとの配色: (表示名, 枠色, 塗り色)
ACTORS = {
    "student": ("児童生徒（区立小中学校）",        RGBColor(0xB9, 0x77, 0x0E), RGBColor(0xFD, 0xEB, 0xD0)),
    "kotaeru": ("こたエール職員",                  RGBColor(0x21, 0x61, 0x8C), RGBColor(0xD6, 0xEA, 0xF8)),
    "police":  ("警察署（練馬区内管轄）",          RGBColor(0x1E, 0x66, 0x41), RGBColor(0xD1, 0xE7, 0xDD)),
    "school":  ("当該校（校長・副校長）",          RGBColor(0x9A, 0x76, 0x0A), RGBColor(0xFC, 0xF3, 0xCF)),
    "shidou":  ("教育指導課 指導主事",             RGBColor(0x6C, 0x34, 0x83), RGBColor(0xE8, 0xDA, 0xEF)),
    "joint":   ("当該校・教育指導課",              RGBColor(0x6C, 0x34, 0x83), RGBColor(0xE8, 0xDA, 0xEF)),
}

SLIDE_W, SLIDE_H = Cm(21.0), Cm(29.7)
LX, LW = 0.85, 9.05    # 左列
RX, RW = 10.8, 9.35    # 右列
LC, RC = LX + LW / 2, RX + RW / 2


def set_text(tf, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True,
             space_after=2):
    """lines: [(text, size, bold, color), ...] 1要素=1段落。\n は段落内改行。"""
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.13)
    tf.margin_top = tf.margin_bottom = Cm(0.06)
    for i, (text, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.02
        p.space_after = Pt(space_after)
        for j, seg in enumerate(text.split("\n")):
            r = p.add_run()
            r.text = seg
            if j > 0:
                br = r._r.makeelement(qn('a:br'), {})
                r._r.addprevious(br)
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = JP_FONT
            rPr = r._r.get_or_add_rPr()
            ea = rPr.makeelement(qn('a:ea'), {})
            ea.set('typeface', JP_FONT)
            rPr.append(ea)


def add_box(slide, shape_type, x, y, w, h, fill, line_color, line_w=1.25):
    sp = slide.shapes.add_shape(shape_type, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def add_arrow(slide, x1, y1, x2, y2, color=LINE_DARK, weight=1.75):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(weight)
    conn.shadow.inherit = False
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)
    return conn


def badge(slide, cx, cy, num):
    d = 0.62
    b = add_box(slide, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, BADGE_FILL, WHITE, 1.25)
    set_text(b.text_frame, [(str(num), 10, True, WHITE)],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False)


def flow_box(slide, num, x, w, y, h, actor, title, body=None, extra=None,
             title_size=10, body_size=9):
    """手順ボックス。extra: 追加段落 [(text, size, bold, color), ...]"""
    name, edge, fill = ACTORS[actor]
    box = add_box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, edge, 1.5)
    box.adjustments[0] = 0.09
    lines = [(f"【{name}】", 8.5, True, edge), (title, title_size, True, TEXT_DARK)]
    if body:
        lines.append((body, body_size, False, TEXT_DARK))
    if extra:
        lines.extend(extra)
    set_text(box.text_frame, lines, anchor=MSO_ANCHOR.MIDDLE, space_after=1)
    if num is not None:
        badge(slide, x - 0.44, y + h / 2, num)
    return box


def col_band(slide, x, w, y, text):
    band = add_box(slide, MSO_SHAPE.RECTANGLE, x, y, w, 0.75, BADGE_FILL, None)
    set_text(band.text_frame, [(text, 10, True, WHITE)],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    return band


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    s = prs.slides.add_slide(prs.slide_layouts[6])

    # ---------------- ヘッダー
    tb = s.shapes.add_textbox(Cm(0.85), Cm(0.25), Cm(19.35), Cm(0.9))
    set_text(tb.text_frame, [("児童生徒の自殺企図に関する相談対応フロー", 15, True, TEXT_DARK)],
             align=PP_ALIGN.CENTER)
    sub = s.shapes.add_textbox(Cm(0.85), Cm(1.0), Cm(19.35), Cm(0.5))
    set_text(sub.text_frame, [("〜こたエール（東京都のネット・スマホトラブル相談窓口）にメールで相談が入った場合〜", 9, False, TEXT_DARK)],
             align=PP_ALIGN.CENTER)

    pr = add_box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.85, 1.6, 19.35, 0.75,
                 RGBColor(0xEA, 0xF0, 0xF8), BADGE_FILL, 1.0)
    pr.adjustments[0] = 0.16
    set_text(pr.text_frame, [
        ("対応の3原則　① 最優先で即時対応　② 一人で抱え込まず組織で対応　③ 時系列で記録（日時・相手・内容）", 9.5, True, BADGE_FILL)],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    wn = add_box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.85, 2.5, 19.35, 0.75, WARN_FILL, WARN_EDGE, 1.25)
    wn.adjustments[0] = 0.16
    set_text(wn.text_frame, [
        ("⚠ 切迫した危険（今まさに実行するおそれ等）が読み取れる場合は、フローと並行して直ちに110番通報", 9.5, True, WARN_EDGE)],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # ---------------- 左列（フェーズ1）
    col_band(s, LX, LW, 3.55, "フェーズ1　覚知 〜 児童生徒の特定（覚知後 直ちに）")

    y = 4.6
    flow_box(s, 1, LX, LW, y, 1.9, "student",
             "こたエールに自殺企図に関する相談メールを送信",
             "区が貸与している学習用アカウント（メールアドレス）から送信される。")
    add_arrow(s, LC, y + 1.9, LC, y + 2.3)

    y = 6.85
    flow_box(s, 2, LX, LW, y, 3.0, "kotaeru",
             "教育指導課の指導主事へ電話で情報提供",
             "※区貸与アカウントは外部からの受信が制限され、本人に返信できないため。",
             extra=[("聞き取り：□アドレス □受信日時 □内容（方法・時期の具体性＝切迫度）□氏名・学校の手がかり □折返し先", 8.5, False, BADGE_FILL)])
    add_arrow(s, LC, y + 3.0, LC, y + 3.35)

    y = 10.2
    flow_box(s, 3, LX, LW, y, 2.6, "shidou",
             "直ちに課内で共有し、時系列記録を開始",
             "課長・統括指導主事へ即報告。対応は必ず複数人で。生命・身体の保護に必要な個人情報の提供・取得は本人同意不要（個人情報保護法27条1項2号）。")
    add_arrow(s, LC, y + 2.6, LC, y + 2.95)

    # 同時進行ブロック
    gy = 13.15
    grp = add_box(s, MSO_SHAPE.ROUNDED_RECTANGLE, LX - 0.12, gy, LW + 0.24, 5.65, None, ORANGE, 1.25)
    grp.line.dash_style = MSO_LINE.DASH
    grp.adjustments[0] = 0.05
    tag = add_box(s, MSO_SHAPE.ROUNDED_RECTANGLE, LX + 0.25, gy - 0.3, 2.5, 0.6, WHITE, ORANGE, 1.25)
    set_text(tag.text_frame, [("【同時進行】", 9, True, ORANGE)],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False)

    flow_box(s, 4, LX + 0.25, LW - 0.5, gy + 0.45, 2.3, "shidou",
             "教育ICT環境整備係（教育施設課）へ児童生徒の照会を依頼",
             "メールアドレスを伝達。区貸与アカウントのため在籍校・学年・氏名を特定できる。")
    flow_box(s, 5, LX + 0.25, LW - 0.5, gy + 3.0, 2.3, "police",
             "警察署から指導主事へ、当該児童生徒の照会の連絡が入る",
             "こたエールは警察にも情報提供している。担当係・担当者名・折返し先を必ず控える。")

    # 判断（ひし形）
    dy = 19.15
    dw, dh = 7.6, 2.0
    dx = LC - dw / 2
    add_arrow(s, LC, gy + 5.65, LC, dy)
    dm = add_box(s, MSO_SHAPE.DIAMOND, dx, dy, dw, dh, WHITE, LINE_DARK, 1.5)
    set_text(dm.text_frame, [("児童生徒を\n特定できたか", 10, True, TEXT_DARK)],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # いいえ → 待機（右下に配置）
    nbx, nbw = LX + 4.75, LW - 4.75
    nb = add_box(s, MSO_SHAPE.ROUNDED_RECTANGLE, nbx, dy + 2.35, nbw, 1.6, NOTE_FILL, LINE_DARK, 1.0)
    nb.adjustments[0] = 0.1
    set_text(nb.text_frame, [
        ("照会結果を待つ", 9.5, True, TEXT_DARK),
        ("判明し次第、直ちに警察の担当者へ連絡する。", 8.5, False, TEXT_DARK)], anchor=MSO_ANCHOR.MIDDLE)
    add_arrow(s, dx + dw, dy + dh / 2, nbx + nbw / 2, dy + 2.35)
    lb = s.shapes.add_textbox(Cm(dx + dw - 0.15), Cm(dy + 1.5), Cm(1.8), Cm(0.5))
    set_text(lb.text_frame, [("いいえ", 9, True, TEXT_DARK)], wrap=False)

    # はい → 右列へ
    cy = 23.35
    add_arrow(s, LC - 1.5, dy + dh, LC - 1.5, cy)
    hb = s.shapes.add_textbox(Cm(LC - 3.0), Cm(dy + dh + 0.4), Cm(1.4), Cm(0.5))
    set_text(hb.text_frame, [("はい", 9, True, TEXT_DARK)], align=PP_ALIGN.RIGHT, wrap=False)
    add_arrow(s, nbx + nbw / 2, dy + 3.95, nbx + nbw / 2, cy - 0.05)
    cv = add_box(s, MSO_SHAPE.PENTAGON, LX, cy, LW, 0.95, LINE_DARK, None)
    set_text(cv.text_frame, [("特定できたら → 右の列(フェーズ2)へ", 10.5, True, WHITE)],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # ---------------- 右列（フェーズ2〜3）
    col_band(s, RX, RW, 3.55, "フェーズ2　警察との連携 〜 安否確認（特定後 直ちに）")
    add_arrow(s, RC, 4.3, RC, 4.65)

    y = 4.65
    flow_box(s, 6, RX, RW, y, 3.0, "shidou",
             "警察へ ①学校名 ②学年 ③氏名 を伝える",
             "警察が当該校へ連絡する前に、指導主事から当該校へ事前連絡する旨を伝え、当該校に何を伝えるべきか助言をもらう。",
             extra=[("助言の例：事前に保護者に連絡しておく／保護者からの虐待が疑われるため、警察が連絡するまでは何もしない　等", 8.5, False, ORANGE)])
    add_arrow(s, RC, y + 3.0, RC, y + 3.35)

    y = 8.0
    flow_box(s, 7, RX, RW, y, 3.3, "shidou",
             "当該校の校長（不在時は副校長）へ電話で事前連絡",
             "・警察から在籍確認の連絡が入ること\n・警察の助言に基づく対応方法（保護者への連絡の要否 等）\n・校内で知る教職員は必要最小限にとどめること",
             extra=[("※虐待疑い時は保護者に連絡せず、子供家庭支援センター・児童相談所へ通告（児童虐待防止法6条）", 8.5, False, WARN_EDGE)])
    add_arrow(s, RC, y + 3.3, RC, y + 3.65)

    y = 11.65
    flow_box(s, 8, RX, RW, y, 1.9, "police",
             "当該校へ連絡し、在籍確認",
             "在籍していれば ①氏名 ②保護者の連絡先 ③住所 を確認する。")
    add_arrow(s, RC, y + 1.9, RC, y + 2.25)

    y = 13.9
    flow_box(s, 9, RX, RW, y, 2.3, "police",
             "当該児童生徒の自宅を訪問し、安否確認",
             "必要に応じ保護等の措置（警察官職務執行法3条）。完了まで指導主事・当該校とも電話を受けられる体制を維持する。")
    add_arrow(s, RC, y + 2.3, RC, y + 2.65)

    y = 16.55
    flow_box(s, 10, RX, RW, y, 1.4, "police",
             "安否確認の結果を当該校へ連絡")
    add_arrow(s, RC, y + 1.4, RC, y + 1.75)

    y = 18.3
    flow_box(s, 11, RX, RW, y, 2.2, "school",
             "警察からの情報を教育指導課へ報告",
             "①安否・現在の状況　②警察がとった措置　③保護者の状況\n④学校が把握している本人の様子（欠席・友人関係・いじめの有無等）")
    add_arrow(s, RC, y + 2.2, RC, y + 2.55)

    y = 20.85
    col_band(s, RX, RW, y, "フェーズ3　事後の支援（翌日以降・継続）")
    flow_box(s, 12, RX, RW, y + 0.75, 2.6, "joint",
             "校内支援体制を立ち上げ、継続的に支援",
             "校内委員会（管理職・担任・養護教諭・SC・SSW）で支援方針を決定。保護者と連携し、必要に応じ医療機関へ接続。本人にはTALKの原則で関わり、経過を教育指導課へ定期報告。")

    # ---------------- 根拠・参考資料
    ry = 24.6
    rb = add_box(s, MSO_SHAPE.RECTANGLE, 0.85, ry, 19.35, 3.3, WHITE, BADGE_FILL, 1.0)
    set_text(rb.text_frame, [
        ("【根拠・参考資料】", 9, True, BADGE_FILL),
        ("・文部科学省「教師が知っておきたい子どもの自殺予防」（平成21年3月）…組織的対応・記録・TALKの原則（Tell 心配を言葉で伝える／Ask 死にたい気持ちを率直に尋ねる／Listen 傾聴する／Keep safe 一人にせず安全を確保）", 8, False, TEXT_DARK),
        ("・文部科学省「子供に伝えたい自殺予防」（平成26年7月）／自殺総合対策大綱（令和4年10月）／こどもの自殺対策緊急強化プラン（令和5年6月）／自殺対策基本法", 8, False, TEXT_DARK),
        ("・個人情報保護法27条1項2号（生命・身体の保護に必要な場合は本人同意なく第三者提供可）／児童虐待防止法6条（通告義務）／警察官職務執行法3条（保護）", 8, False, TEXT_DARK),
        ("※電話でのやり取りは復唱して確認し、直後に記録する。情報は「知る必要のある者」に限定して共有する。", 8, False, TEXT_DARK)],
        space_after=2)

    foot = s.shapes.add_textbox(Cm(0.85), Cm(28.95), Cm(19.35), Cm(0.55))
    set_text(foot.text_frame, [("練馬区教育委員会 教育指導課", 8, False, NOTE_EDGE)],
             align=PP_ALIGN.RIGHT)

    out = "自殺企図相談対応フロー図.pptx"
    prs.save(out)
    print(f"saved: {out}")


if __name__ == "__main__":
    main()
