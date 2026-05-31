# -*- coding: utf-8 -*-
"""
v10(30枚)→v11(32枚)
1. ペアトーク用スライド2枚を追加（旧S11の後／旧S25の後）
2. 参考文献の誤りを修正：
   - 「第3章第2節2」→「第3章第3節2」（S5/S15/S30の3箇所）
   - 「学習評価の在り方ハンドブック（中学校編）」→「（小・中学校編）」
   - 東京都資料の正式名称を併記
3. 新スライド2枚にも発表者ノート埋め込み
"""
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

WATER_BLUE=RGBColor(0x5B,0xC0,0xDE); LIGHT_WATER=RGBColor(0xD6,0xEE,0xFA)
DEEP_WATER=RGBColor(0x2A,0x96,0xC2); NAVY=RGBColor(0x1F,0x3A,0x6E)
ORANGE_KW=RGBColor(0xEB,0x6C,0x15); YELLOW=RGBColor(0xFF,0xE0,0x66)
WHITE=RGBColor(0xFF,0xFF,0xFF); BLACK=RGBColor(0x00,0x00,0x00); DARK_GRAY=RGBColor(0x40,0x40,0x40)
GREEN=RGBColor(0x07,0xA9,0x73)
F_MAIN="メイリオ"

prs = Presentation('/tmp/output/講義スライド_石神井南中_20260603_v10.pptx')
SW,SH = prs.slide_width, prs.slide_height
assert len(prs.slides)==30

def blank(): return prs.slides.add_slide(prs.slide_layouts[6])

def chapter_bar(slide,text,fs=24,color=DEEP_WATER):
    bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,SW,Emu(720000))
    bar.fill.solid(); bar.fill.fore_color.rgb=color; bar.line.fill.background()
    tf=bar.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(300000); tf.margin_right=Emu(100000); tf.margin_top=Emu(20000); tf.margin_bottom=Emu(20000)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT; r=p.add_run(); r.text=text
    r.font.size=Pt(fs); r.font.bold=True; r.font.color.rgb=WHITE; r.font.name=F_MAIN

def sub_heading(slide,text,fs=28):
    bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Emu(830000),SW,Emu(700000))
    bar.fill.solid(); bar.fill.fore_color.rgb=LIGHT_WATER; bar.line.fill.background()
    tf=bar.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(300000); tf.margin_right=Emu(200000); tf.margin_top=Emu(20000); tf.margin_bottom=Emu(20000)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT; r=p.add_run(); r.text=text
    r.font.size=Pt(fs); r.font.bold=True; r.font.color.rgb=NAVY; r.font.name=F_MAIN
    ul=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Emu(1530000),SW,Emu(40000))
    ul.fill.solid(); ul.fill.fore_color.rgb=WATER_BLUE; ul.line.fill.background()

def box(slide,left,top,width,height,lines,default_size=24,default_color=BLACK,line_spacing=1.3,
        align=PP_ALIGN.LEFT,fill_color=WHITE,border_color=WATER_BLUE,border_width=3.0,
        vanchor=MSO_ANCHOR.MIDDLE,default_bold=True):
    b=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,left,top,width,height)
    b.fill.solid(); b.fill.fore_color.rgb=fill_color
    if border_color: b.line.color.rgb=border_color; b.line.width=Pt(border_width)
    else: b.line.fill.background()
    tf=b.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(140000); tf.margin_right=Emu(140000); tf.margin_top=Emu(80000); tf.margin_bottom=Emu(80000)
    tf.vertical_anchor=vanchor
    for i,item in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if isinstance(item,list):
            for j,ri in enumerate(item):
                rt,ro = ri if isinstance(ri,tuple) else (ri,{})
                if j==0: p.alignment=ro.get('align',align); p.line_spacing=ro.get('line_spacing',line_spacing)
                r=p.add_run(); r.text=rt
                r.font.size=Pt(ro.get('size',default_size)); r.font.bold=ro.get('bold',default_bold)
                r.font.color.rgb=ro.get('color',default_color); r.font.name=F_MAIN
        else:
            text,o = item if isinstance(item,tuple) else (item,{})
            p.alignment=o.get('align',align); p.line_spacing=o.get('line_spacing',line_spacing)
            r=p.add_run(); r.text=text
            r.font.size=Pt(o.get('size',default_size)); r.font.bold=o.get('bold',default_bold)
            r.font.color.rgb=o.get('color',default_color); r.font.name=F_MAIN
    return b

# ============================================================
# (修正1) S5の出典バー：「第2節」→「第3節」
# ============================================================
for sh in prs.slides[4].shapes:
    if sh.has_text_frame and '第3章第2節2' in sh.text_frame.text:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if '第3章第2節2' in r.text:
                    r.text = r.text.replace('第3章第2節2', '第3章第3節2')

# ============================================================
# (修正2) S15の出典バー：「第2節」→「第3節」
# ============================================================
for sh in prs.slides[14].shapes:
    if sh.has_text_frame and '第3章第2節2' in sh.text_frame.text:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if '第3章第2節2' in r.text:
                    r.text = r.text.replace('第3章第2節2', '第3章第3節2')

# ============================================================
# (修正3) S30 参考文献：章節番号・ハンドブック名・東京都資料の修正
# ============================================================
s_ref = prs.slides[29]
for sh in s_ref.shapes:
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if '第3章第2節2' in r.text:
                    r.text = r.text.replace('第3章第2節2', '第3章第3節2')
                if '学習評価の在り方ハンドブック（中学校編）' in r.text:
                    r.text = r.text.replace('学習評価の在り方ハンドブック（中学校編）',
                                            '学習評価の在り方ハンドブック（小・中学校編）')
                if r.text == '● 「指導と評価の一体化を目指して」 東京都教育委員会':
                    r.text = '● 子供たちに未来の創り手となるために必要な資質・能力を育む'
# 東京都資料：別行を追加するために、該当箇所の構造を変更
for sh in s_ref.shapes:
    if sh.has_text_frame and '子供たちに未来の創り手' in sh.text_frame.text:
        # 「子供たちに未来…」のparaの直後に副題を追加
        from lxml import etree
        target_p = None
        for p in sh.text_frame.paragraphs:
            if any('子供たちに未来の創り手' in r.text for r in p.runs):
                target_p = p
                break
        if target_p is not None:
            new_p_el = etree.SubElement(sh.text_frame._txBody, '{http://schemas.openxmlformats.org/drawingml/2006/main}p')
            target_p._p.addnext(new_p_el)
            from pptx.text.text import _Paragraph
            new_p = _Paragraph(new_p_el, target_p._parent)
            new_p.line_spacing = 1.25
            r = new_p.add_run()
            r.text = '　 指導と評価の一体化を目指して　Ⅰ理論編／Ⅱ実践編　東京都教育委員会（令和2年9月）'
            r.font.size = Pt(13); r.font.color.rgb = DARK_GRAY; r.font.bold = False; r.font.name = F_MAIN

# ============================================================
# (追加1) ペアトーク① ─ S11(授業改善と評価は一体)の後に挿入予定
# ============================================================
p1 = blank()
chapter_bar(p1, "【ペアトーク①】　─　ここで一息、隣の方と", color=GREEN)
sub_heading(p1, "2分で話してみてください")

# 中央の大問
box(p1, Emu(400000), Emu(1900000), Emu(11400000), Emu(1600000),
    [[("最近の授業で、", {'size':28,'color':BLACK,'align':PP_ALIGN.CENTER})],
     [("「主体的な学び」「対話的な学び」「深い学び」", {'size':30,'color':ORANGE_KW,'align':PP_ALIGN.CENTER})],
     [("が見えた場面はありましたか？", {'size':28,'color':BLACK,'align':PP_ALIGN.CENTER})]],
    fill_color=LIGHT_WATER, border_width=5.0, border_color=GREEN, line_spacing=1.3)

# 補助メッセージ
box(p1, Emu(400000), Emu(3700000), Emu(11400000), Emu(1100000),
    [[("○ どんな小さなことでもOK", {'size':22,'color':NAVY})],
     [("○ 教科特性を超えて、一場面を思い出してみる", {'size':22,'color':NAVY})],
     [("○ 「逆に見えなかった…」もOKです", {'size':22,'color':DARK_GRAY})]],
    border_width=3.0, line_spacing=1.5)

# 下部のタイマー＆方法
box(p1, Emu(400000), Emu(5000000), Emu(5400000), Emu(900000),
    [("⏱ 2分", {'size':40,'bold':True,'color':WHITE,'align':PP_ALIGN.CENTER})],
    fill_color=ORANGE_KW, border_color=ORANGE_KW, border_width=2.0)
box(p1, Emu(6400000), Emu(5000000), Emu(5400000), Emu(900000),
    [("隣の方とペアで", {'size':24,'bold':True,'color':NAVY,'align':PP_ALIGN.CENTER})],
    fill_color=YELLOW, border_color=ORANGE_KW, border_width=2.0)

box(p1, Emu(400000), Emu(6010000), Emu(11400000), Emu(360000),
    [("→ このあと、「主体的に学習に取り組む態度」の評価をどう見取るか、お話します",
      {'size':14,'color':DARK_GRAY,'align':PP_ALIGN.CENTER})],
    fill_color=WHITE, border_color=None)

# ============================================================
# (追加2) ペアトーク② ─ S25(教科会で揃える3つ)の後に挿入予定
# ============================================================
p2 = blank()
chapter_bar(p2, "【ペアトーク②】　─　自分の教科で「Ｂの姿」を描く", color=GREEN)
sub_heading(p2, "3分（個人考1分→ペアでシェア2分）")

# 中央の大問
box(p2, Emu(400000), Emu(1900000), Emu(11400000), Emu(1600000),
    [[("あなたの教科で、", {'size':28,'color':BLACK,'align':PP_ALIGN.CENTER})],
     [("次の単元の", {'size':28,'color':BLACK,'align':PP_ALIGN.CENTER}),
      ("「Ｂの姿」", {'size':32,'color':ORANGE_KW,'align':PP_ALIGN.CENTER}),
      ("を1つ、言葉にしてみよう", {'size':28,'color':BLACK,'align':PP_ALIGN.CENTER})]],
    fill_color=LIGHT_WATER, border_width=5.0, border_color=GREEN, line_spacing=1.3)

# 例の提示
box(p2, Emu(400000), Emu(3700000), Emu(11400000), Emu(1100000),
    [[("○ ", {'size':22,'color':NAVY}),
      ("「〜できている」「〜と書ける」「〜を選んでいる」", {'size':22,'color':NAVY})],
     [("○ 例（数学）：", {'size':20,'color':NAVY}),
      ("「式の操作の途中で、自分のミスに気づいて修正できている」", {'size':20,'color':NAVY})],
     [("○ うまく書けなくてもOK ─ 言葉にしようとすることが大事", {'size':20,'color':DARK_GRAY})]],
    border_width=3.0, line_spacing=1.4)

# 下部のタイマー＆方法
box(p2, Emu(400000), Emu(5000000), Emu(5400000), Emu(900000),
    [("⏱ 3分", {'size':40,'bold':True,'color':WHITE,'align':PP_ALIGN.CENTER})],
    fill_color=ORANGE_KW, border_color=ORANGE_KW, border_width=2.0)
box(p2, Emu(6400000), Emu(5000000), Emu(5400000), Emu(900000),
    [[("個人考1分→ペアでシェア2分", {'size':20,'bold':True,'color':NAVY,'align':PP_ALIGN.CENTER})]],
    fill_color=YELLOW, border_color=ORANGE_KW, border_width=2.0)

box(p2, Emu(400000), Emu(6010000), Emu(11400000), Emu(360000),
    [("→ ここで描いた「Ｂの姿」を、このあとのグループ協議で持ち寄ります",
      {'size':14,'color':DARK_GRAY,'align':PP_ALIGN.CENTER})],
    fill_color=WHITE, border_color=None)

# ============================================================
# 並べ替え：v10の30枚 + 新規2枚 = 32枚
# v10構成：0..29 → S1..S30
# 新規：30=ペアトーク①, 31=ペアトーク②
# 望ましい順：[0..10, 30, 11..24, 31, 25..29]
# ============================================================
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
desired = list(range(11)) + [30] + list(range(11,25)) + [31] + list(range(25,30))
assert sorted(desired) == list(range(32)), desired
for el in ids: sldIdLst.remove(el)
for i in desired: sldIdLst.append(ids[i])

# ============================================================
# 新2枚にノート埋め込み
# ============================================================
notes_pair1 = """【ペアトーク① ─ 自分の授業に引き寄せる】 ─ 目安2分

ここで一度、皆さんに話していただく時間を取ります。私の話を聞き続けてばかりだと、自分の授業との接続が見えにくくなります。

隣の先生と、最近のご自分の授業を振り返って、お話ください。

問いは ── 「最近の授業で、『主体的な学び』『対話的な学び』『深い学び』が見えた場面はありましたか？」

どんな小さな場面でもOKです。「あの生徒が、自分から振り返りを書いていた」「あのグループの話し合いが弾んでいた」── そんなレベルで結構です。

逆に「見えなかった…」というのも立派な気づきです。それが、このあとお話する評価の話につながります。

では、2分。お願いします。

※ 講師は時間を計りつつ、教室内を巡回して様子を見る。"""

notes_pair2 = """【ペアトーク② ─ Ｂの姿を試しに描いてみる】 ─ 目安3分

「教科会で揃える3つ」のお話をしました。①Ｂの姿／②評価場面・方法／③総括の方法。

特に最初の「Ｂの姿」 ── これを、ここで一度試しに描いてみていただきたいと思います。このあとのグループ協議の助走です。

問いは ── 「あなたの教科で、次の単元の『Ｂの姿』を1つ、言葉にしてみよう」。

例えば数学なら「式の操作の途中で、自分のミスに気づいて修正できている」。「〜できている」「〜と書ける」「〜を選んでいる」といった、生徒の具体的な姿として描いてみてください。

最初の1分は個人で考えて、メモ用紙やノートに書いてみる。残り2分でペアでシェアしてください。

うまく書けなくても大丈夫です。「言葉にしようとすること」自体が、評価規準を体得していく一番の方法です。

では、お願いします。

※ ここで描いてもらった「Ｂの姿」が、このあとのグループ協議で持ち寄る評価資料の核になる。"""

for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame and '【ペアトーク①】' in sh.text_frame.text:
            s.notes_slide.notes_text_frame.text = notes_pair1
        if sh.has_text_frame and '【ペアトーク②】' in sh.text_frame.text:
            s.notes_slide.notes_text_frame.text = notes_pair2

# フォント統一
from lxml import etree
NS_A='http://schemas.openxmlformats.org/drawingml/2006/main'
def force(run):
    rPr=run._r.find(f'{{{NS_A}}}rPr')
    if rPr is None: rPr=etree.SubElement(run._r,f'{{{NS_A}}}rPr'); run._r.insert(0,rPr)
    for t in ['latin','ea','cs']:
        e=rPr.find(f'{{{NS_A}}}{t}')
        if e is None: e=etree.SubElement(rPr,f'{{{NS_A}}}{t}')
        e.set('typeface','メイリオ')
def wf(shapes):
    n=0
    for sh in shapes:
        if sh.shape_type==MSO_SHAPE_TYPE.GROUP: n+=wf(sh.shapes)
        elif sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs: force(r); n+=1
    return n
print('font runs:', sum(wf(s.shapes) for s in prs.slides))

import os; os.makedirs('/tmp/output',exist_ok=True)
OUT='/tmp/output/講義スライド_石神井南中_20260603_v11.pptx'
prs.save(OUT)
print('saved:', OUT, 'slides:', len(prs.slides))
