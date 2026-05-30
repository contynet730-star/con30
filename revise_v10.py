# -*- coding: utf-8 -*-
"""
v9(28枚)に「今後に向けて」2枚を追加してv10(30枚)を作成
S28: 次期学習指導要領の方向性 ─ 令和7年「論点整理」
S29: 今日の取り組みは次期でこそ生きる
旧S28(参考文献)→S30へ移し、論点整理を追加。新規2枚にもノート埋め込み。
"""
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

WATER_BLUE=RGBColor(0x5B,0xC0,0xDE); LIGHT_WATER=RGBColor(0xD6,0xEE,0xFA)
DEEP_WATER=RGBColor(0x2A,0x96,0xC2); NAVY=RGBColor(0x1F,0x3A,0x6E)
ORANGE_KW=RGBColor(0xEB,0x6C,0x15); RED=RGBColor(0xC0,0x40,0x40)
GREEN=RGBColor(0x07,0xA9,0x73); YELLOW=RGBColor(0xFF,0xE0,0x66)
WHITE=RGBColor(0xFF,0xFF,0xFF); BLACK=RGBColor(0x00,0x00,0x00); DARK_GRAY=RGBColor(0x40,0x40,0x40)
F_MAIN="メイリオ"

prs = Presentation('/tmp/output/講義スライド_発表者ノート付_v9.pptx')
SW,SH = prs.slide_width, prs.slide_height
assert len(prs.slides)==28

def blank(): return prs.slides.add_slide(prs.slide_layouts[6])
def chapter_bar(slide,text,fs=22):
    bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,SW,Emu(720000))
    bar.fill.solid(); bar.fill.fore_color.rgb=DEEP_WATER; bar.line.fill.background()
    tf=bar.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(300000); tf.margin_right=Emu(100000); tf.margin_top=Emu(20000); tf.margin_bottom=Emu(20000)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT; r=p.add_run(); r.text=text
    r.font.size=Pt(fs); r.font.bold=True; r.font.color.rgb=WHITE; r.font.name=F_MAIN
def sub_heading(slide,text,fs=28):
    bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Emu(830000),SW,Emu(700000))
    bar.fill.solid(); bar.fill.fore_color.rgb=LIGHT_WATER; bar.line.fill.background()
    tf=bar.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(300000); tf.margin_right=Emu(200000); tf.margin_top=Emu(20000); tf.margin_bottom=Emu(20000)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT; r=p.add_run(); r.text=text
    r.font.size=Pt(fs); r.font.bold=True; r.font.color.rgb=NAVY; r.font.name=F_MAIN
    ul=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Emu(1530000),SW,Emu(40000))
    ul.fill.solid(); ul.fill.fore_color.rgb=WATER_BLUE; ul.line.fill.background()
def box(slide,left,top,width,height,lines,default_size=24,default_color=BLACK,line_spacing=1.3,
        align=PP_ALIGN.LEFT,fill_color=WHITE,border_color=WATER_BLUE,border_width=3.0,
        vanchor=MSO_ANCHOR.MIDDLE,default_bold=True):
    b=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,left,top,width,height)
    b.fill.solid(); b.fill.fore_color.rgb=fill_color
    if border_color: b.line.color.rgb=border_color; b.line.width=Pt(border_width)
    else: b.line.fill.background()
    tf=b.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(140000); tf.margin_right=Emu(140000); tf.margin_top=Emu(80000); tf.margin_bottom=Emu(80000)
    tf.vertical_anchor=vanchor
    for i,item in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if isinstance(item,list):
            for j,ri in enumerate(item):
                rt,ro = ri if isinstance(ri,tuple) else (ri,{})
                if j==0: p.alignment=ro.get('align',align); p.line_spacing=ro.get('line_spacing',line_spacing)
                r=p.add_run(); r.text=rt
                r.font.size=Pt(ro.get('size',default_size)); r.font.bold=ro.get('bold',default_bold)
                r.font.color.rgb=ro.get('color',default_color); r.font.name=F_MAIN
        else:
            text,o = item if isinstance(item,tuple) else (item,{})
            p.alignment=o.get('align',align); p.line_spacing=o.get('line_spacing',line_spacing)
            r=p.add_run(); r.text=text
            r.font.size=Pt(o.get('size',default_size)); r.font.bold=o.get('bold',default_bold)
            r.font.color.rgb=o.get('color',default_color); r.font.name=F_MAIN
    return b
def add_source_line(slide, text, top=Emu(6000000)):
    tb = slide.shapes.add_textbox(Emu(300000), top, Emu(11500000), Emu(330000))
    tf = tb.text_frame; tf.word_wrap=True
    tf.margin_left=Emu(40000); tf.margin_right=Emu(40000); tf.margin_top=Emu(10000); tf.margin_bottom=Emu(10000)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT
    r=p.add_run(); r.text='出典：'+text
    r.font.size=Pt(13); r.font.color.rgb=DARK_GRAY; r.font.name=F_MAIN

# ============================================================
# NEW1: 今後に向けて(1) ─ 次期学習指導要領の方向性
# ============================================================
n1 = blank()
chapter_bar(n1, "今後に向けて")
sub_heading(n1, "次期学習指導要領の方向性 ─ 令和7年「論点整理」")

# 上部の注意書き
box(n1, Emu(400000), Emu(1820000), Emu(11400000), Emu(620000),
    [[("ここから先は", {'size':18,'color':BLACK}),
      ("「方向性」", {'size':22,'color':ORANGE_KW}),
      ("です。現時点での運用は ", {'size':18,'color':BLACK}),
      ("現行3観点ABC", {'size':22,'color':NAVY}),
      (" のままです", {'size':18,'color':BLACK})]],
    fill_color=YELLOW, border_color=ORANGE_KW, border_width=2.5, align=PP_ALIGN.CENTER)

# 3つのポイント
points = [
    ("①", "評定の対象外へ",
     "「主体的に学習に取り組む態度」を A・B・C の目標準拠評価から外し、\n個人の成長や良さを記述する「個人内評価」へ転換"),
    ("②", "総合所見＋「○」付記",
     "「学びに向かう力」を総合所見欄での個人内評価と、\n「思考・判断・表現」観点別評価への「○」付記の組合せで見る方向"),
    ("③", "形式的評価からの脱却",
     "ノート提出頻度・挙手回数など「勤勉さ」ではなく、\n見通し・粘り強さ・振り返りの「学びの過程」そのものを評価"),
]
y = 2550000
for i,(num,kw,desc) in enumerate(points):
    yy = Emu(y + i*1100000)
    box(n1, Emu(400000), yy, Emu(900000), Emu(950000),
        [(num, {'size':42,'bold':True,'color':WHITE,'align':PP_ALIGN.CENTER})],
        fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
    box(n1, Emu(1400000), yy, Emu(10400000), Emu(950000),
        [(kw, {'size':22,'color':ORANGE_KW}),
         (desc, {'size':16,'color':NAVY})],
        line_spacing=1.2, border_width=3.0)

add_source_line(n1, "中央教育審議会 教育課程企画特別部会「論点整理」令和7年9月25日／検討資料⑧令和8年3月30日")

# ============================================================
# NEW2: 今後に向けて(2) ─ 今日の取り組みは次期でこそ生きる
# ============================================================
n2 = blank()
chapter_bar(n2, "今後に向けて")
sub_heading(n2, "今日の取り組みは「次期」でこそ生きる")

box(n2, Emu(400000), Emu(1820000), Emu(11400000), Emu(650000),
    [[("「今日の研修が無駄になるのでは？」", {'size':20,'color':DARK_GRAY}),
      ("　─　いいえ、むしろ", {'size':20,'color':BLACK}),
      ("強化されます", {'size':22,'color':ORANGE_KW})]],
    fill_color=LIGHT_WATER, border_width=3.0, align=PP_ALIGN.CENTER)

# 3つの「生きる」
keeps = [
    ("Ｂの姿の言語化", "→ 個人内評価の根拠記述に直結",
     "規準が言語化できていれば、所見欄での質の高い記述が書ける"),
    ("形成的評価\n（みとる→気づく→変える→伸びる）", "→ 次期の中心となる視点",
     "「学びの過程」を見る目こそ、次期で最も求められる教師の力"),
    ("教科会で揃える3つ", "→ 観点が減っても価値は不変",
     "個人内評価でも「教科として何を見るか」を揃える力は同じく必要"),
]
y = 2570000
for i,(kw,arrow,desc) in enumerate(keeps):
    yy = Emu(y + i*1000000)
    box(n2, Emu(400000), yy, Emu(4200000), Emu(880000),
        [(kw, {'size':18,'color':WHITE,'align':PP_ALIGN.CENTER})],
        fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0,
        line_spacing=1.15)
    box(n2, Emu(4700000), yy, Emu(7100000), Emu(880000),
        [(arrow, {'size':18,'color':ORANGE_KW}),
         (desc, {'size':15,'color':NAVY})],
        line_spacing=1.2, border_width=3.0)

box(n2, Emu(400000), Emu(5700000), Emu(11400000), Emu(550000),
    [[("方向性が変わっても、", {'size':20,'color':BLACK}),
      ("「学びの過程を見る目」", {'size':22,'color':ORANGE_KW}),
      ("は変わらない", {'size':20,'color':BLACK})]],
    fill_color=YELLOW, border_color=ORANGE_KW, border_width=3.0, align=PP_ALIGN.CENTER)

# ============================================================
# 旧S28(参考文献)に論点整理を追加
# ============================================================
s_ref = prs.slides[27]  # 旧S28（v9で28番目）
# 既存の参考文献ボックスを探して、テキストを差し替え
for sh in s_ref.shapes:
    if sh.has_text_frame and '指導要領' in sh.text_frame.text and '解説' in sh.text_frame.text:
        # この箱が参考文献本体
        tf = sh.text_frame
        # 一旦全消去
        # 既存のparagraphsを保持しつつ追加するのが安全（最初のpを除いて削除→再構築）
        # シンプルに上書き
        from copy import deepcopy
        # text_frame全消去
        for p in list(tf.paragraphs[1:]):
            p._p.getparent().remove(p._p)
        tf.paragraphs[0].clear()

        refs = [
            ("● 中央教育審議会 教育課程企画特別部会「論点整理」（令和7年9月25日）", ORANGE_KW, True),
            ("　 教育課程部会 総則・評価特別部会「学習評価の在り方について」", DARK_GRAY, False),
            ("　 検討資料⑧（令和8年3月30日）", DARK_GRAY, False),
            ("● 中学校学習指導要領 解説　総則編　文部科学省", NAVY, True),
            ("　 第3章第2節2「学習評価の充実」", DARK_GRAY, False),
            ("● 学習評価の在り方ハンドブック（中学校編）", NAVY, True),
            ("　 文部科学省 国立教育政策研究所教育課程研究センター", DARK_GRAY, False),
            ("● 「指導と評価の一体化」のための学習評価に関する参考資料", NAVY, True),
            ("　 文部科学省 国立教育政策研究所教育課程研究センター", DARK_GRAY, False),
            ("● 「指導と評価の一体化を目指して」 東京都教育委員会", NAVY, True),
            ("● 児童生徒の学習評価の在り方について（報告）", NAVY, True),
            ("　 中央教育審議会 平成31年1月21日", DARK_GRAY, False),
        ]
        from pptx.util import Pt as _Pt
        for i,(t,c,bold) in enumerate(refs):
            p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.line_spacing = 1.25
            r = p.add_run(); r.text = t
            r.font.size = _Pt(16 if bold else 13)
            r.font.color.rgb = c
            r.font.bold = bold
            r.font.name = F_MAIN
        break

# ============================================================
# 並べ替え：v9の28枚 + 新規2枚 = 30枚
# v9の構成: 0..26 = S1..S27, 27 = S28(参考文献)
# 新規: 28 = 今後に向けて(1), 29 = 今後に向けて(2)
# 望ましい順: [0..26, 28, 29, 27]
# ============================================================
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)  # 30
desired = list(range(27)) + [28, 29, 27]
assert sorted(desired) == list(range(30))
for el in ids: sldIdLst.remove(el)
for i in desired: sldIdLst.append(ids[i])

# ============================================================
# 新2枚にノート埋め込み
# ============================================================
notes_new = {
    27: """【今後に向けて(1) ─ 次期学習指導要領の方向性】 ─ 目安1.5分

ここからは、研修の補足として、文部科学省の最新の動きをお伝えします。

令和7年9月25日、中央教育審議会の教育課程企画特別部会から「論点整理」が出されました。次期学習指導要領の方向性が示されています。

まず最初にお伝えしておきたいのは ── ここから話す内容は「方向性」です。現時点（令和8年6月）での皆様の評価運用は、これまでお話してきた現行の3観点ABCのままで結構です。

その上で、次期で示された大きな方向は3つです。

① 「主体的に学習に取り組む態度」を、A・B・Cの目標準拠評価から外し、個人内評価へ転換する方向。
② 「学びに向かう力」を、総合所見欄での個人内評価＋「思考・判断・表現」の観点に「○」を付記する組み合わせで見る方向。
③ ノート提出頻度や挙手回数など「勤勉さ」での評価から脱却し、「学びの過程」そのものを見る方向。

特に③は、今日繰り返しお話してきた「形成的評価」「学びの姿を見取る」という発想と完全に一致しています。""",

    28: """【今後に向けて(2) ─ 今日の取り組みは次期でこそ生きる】 ─ 目安1.5分

「今日聞いた内容、次期で評価方法が変わるなら、無駄になるんじゃないの？」

そう思われる先生もいらっしゃるかもしれません。私自身もそこは大事な論点だと思います。結論から言うと、無駄になるどころか、むしろ「次期でこそ生きる」内容です。

具体に3点。

1点目、「Ｂの姿の言語化」。これは次期では「個人内評価の根拠記述」にそのまま使えます。所見欄に質の高い記述を書くには、まず「何を見ているか」を言語化できていることが必須です。

2点目、形成的評価のサイクル ── みとる→気づく→変える→伸びる。これは次期学習評価が最も重視する「学びの過程を見る目」そのものです。

3点目、教科会で揃える力。観点が減って個人内評価になっても、「教科として何を見るか」を揃える教科会の力は、同じく ── あるいは今以上に ── 必要になります。

方向性は変わっても、「学びの過程を見る目」は変わらない。今日身につけていただいたものは、必ず次期で生きます。""",
}

for idx_in_orig, txt in notes_new.items():
    slide = prs.slides[idx_in_orig]  # 元のadd_slide順
    # ↑ ただしsldIdLstを並べ替えたので prs.slides は順番が変わっている
    # → 直接ターゲットスライドを取り直す
# Better approach: target by detecting heading text
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame and '次期学習指導要領の方向性' in sh.text_frame.text:
            s.notes_slide.notes_text_frame.text = notes_new[27]
        if sh.has_text_frame and '今日の取り組みは「次期」でこそ生きる' in sh.text_frame.text:
            s.notes_slide.notes_text_frame.text = notes_new[28]

# フォント統一
from lxml import etree
NS_A='http://schemas.openxmlformats.org/drawingml/2006/main'
def force(run):
    rPr=run._r.find(f'{{{NS_A}}}rPr')
    if rPr is None: rPr=etree.SubElement(run._r,f'{{{NS_A}}}rPr'); run._r.insert(0,rPr)
    for t in ['latin','ea','cs']:
        e=rPr.find(f'{{{NS_A}}}{t}')
        if e is None: e=etree.SubElement(rPr,f'{{{NS_A}}}{t}')
        e.set('typeface','メイリオ')
def wf(shapes):
    n=0
    for sh in shapes:
        if sh.shape_type==MSO_SHAPE_TYPE.GROUP: n+=wf(sh.shapes)
        elif sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs: force(r); n+=1
    return n
print('font runs:', sum(wf(s.shapes) for s in prs.slides))

import os; os.makedirs('/tmp/output',exist_ok=True)
OUT='/tmp/output/講義スライド_石神井南中_20260603_v10.pptx'
prs.save(OUT)
print('saved:', OUT, 'slides:', len(prs.slides))
