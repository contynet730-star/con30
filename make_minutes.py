# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

JP_FONT = "游明朝"
JP_GOTHIC = "游ゴシック"

doc = Document()

# 既定スタイル（日本語フォント設定）
style = doc.styles["Normal"]
style.font.name = JP_FONT
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), JP_FONT)

# 余白
for s in doc.sections:
    s.top_margin = Cm(2.0)
    s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(2.2)
    s.right_margin = Cm(2.2)


def set_jp(run, font=JP_FONT):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(18)
    set_jp(r, JP_GOTHIC)
    p.paragraph_format.space_after = Pt(6)
    return p


def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x6E)
    set_jp(r, JP_GOTHIC)
    # 下線（段落罫線）
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1F3B6E")
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def sub(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(10.5)
    set_jp(r, JP_GOTHIC)
    return p


def body(text, indent=0.4, bullet=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(("・" if bullet else "") + text)
    set_jp(r)
    r.font.size = Pt(10.5)
    return p


# ===== タイトル =====
title("議　事　録")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("（個人情報の取扱いに係る保護者面談）")
set_jp(r, JP_GOTHIC)
r.font.size = Pt(11)
p.paragraph_format.space_after = Pt(10)

# ===== 開催概要（表） =====
info = [
    ("日　時", "2026年6月22日（月）　午後"),
    ("場　所", "教育委員会　会議室"),
    ("出席者", "都丸統括指導主事（教育委員会・統括指導主事）\n"
              "紺多指導主事（教育委員会・指導主事／本件担当）\n"
              "保護者（児童・井上日向さんの母）\n"
              "柳沢議員（本面談の設定・同席）"),
    ("議　題", "学校が保護者の同意なく児童の個人情報（在籍・クラス・下校時刻）を放課後等デイサービス事業者へ提供した件について"),
]
table = doc.add_table(rows=len(info), cols=2)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.columns[0].width = Cm(2.8)
table.columns[1].width = Cm(14.0)
for i, (k, v) in enumerate(info):
    c0 = table.rows[i].cells[0]
    c1 = table.rows[i].cells[1]
    c0.width = Cm(2.8)
    c1.width = Cm(14.0)
    # shade key cell
    tcPr = c0._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "E8EDF5")
    tcPr.append(shd)
    pr0 = c0.paragraphs[0]
    run0 = pr0.add_run(k)
    run0.bold = True
    set_jp(run0, JP_GOTHIC)
    run0.font.size = Pt(10.5)
    lines = v.split("\n")
    pr1 = c1.paragraphs[0]
    for j, ln in enumerate(lines):
        if j > 0:
            pr1 = c1.add_paragraph()
        rr = pr1.add_run(ln)
        set_jp(rr)
        rr.font.size = Pt(10.5)

doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ===== 1. 面談の趣旨 =====
heading("1. 面談の趣旨")
body("学校および教育委員会の対応により保護者に不安・心配を与えたことについて、教育委員会（都丸統括指導主事）より冒頭で謝罪があった。")
body("本面談は、保護者の心情・経緯を改めて確認したうえで、教育委員会として今後の対応方針を共有することを目的として、柳沢議員の調整により設定された。")

# ===== 2. 経緯（時系列） =====
heading("2. 経緯（時系列）")
rows = [
    ("4月30日", "放課後等デイサービス事業者「日だまりパレット」から学校へ連絡があり、学校（副校長）が保護者の同意を得ないまま、児童（2年1組・井上日向さん）の下校時刻等を事業者へ伝えた。なお保護者は当該事業者への通所を未決定であり、5月1日に保護者から事業者へ連絡する予定であった。"),
    ("4月30日", "同日、事業者は保護者へ架電したが連絡がつかず、その後学校へ連絡した（事業者側の説明）。保護者は当日夕方、事業者に確認し経緯の説明を受けた。"),
    ("5月1日（朝）", "保護者（夫）が学校へ架電。副校長が対応し、この電話口では「第三者（事業者）が本物かどうかは確認していない」旨を説明した。保護者は経緯をまとめた説明を求めたが、短時間で終了。"),
    ("5月1日（夜）", "副校長から架電。説明のないまま事例の話を始め、「自分の話は以上」とし、校長への取次ぎに応じなかった。"),
    ("ゴールデンウィーク中", "学校からの対応・連絡なし。"),
    ("5月7日", "担任との15分間の面談。担任より「保護者から聞いていないのに勝手に伝えてしまった」と謝罪があり、木曜日の時間割（下校時刻）を伝えた旨の説明があった。"),
    ("5月上旬", "保護者が教育委員会（加藤氏）に校長との面談を依頼。後日、校長から保護者へ架電があり、「今回の件はミスではない」「担任が掛け直しているので安全だった」「白黒つける案件ではない」等の説明があった。保護者が納得できない旨を伝えると、校長は「持ち帰る」と回答。"),
    ("その後", "校長から再度架電・直接面談。電話と同一内容の説明が繰り返された。"),
    ("5月18日", "保護者が「要望書」を提出（回答期限を5月29日に設定）。"),
    ("5月29日（前後）", "学校（校長名）より「回答書」が提出された（出張等を理由とするやり取りを経て交付）。"),
]
t = doc.add_table(rows=len(rows) + 1, cols=2)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = t.rows[0].cells
for idx, h in enumerate(["時　期", "内　容"]):
    tcPr = hdr[idx]._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "1F3B6E")
    tcPr.append(shd)
    rr = hdr[idx].paragraphs[0].add_run(h)
    rr.bold = True; rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_jp(rr, JP_GOTHIC); rr.font.size = Pt(10.5)
for i, (d, c) in enumerate(rows, start=1):
    cells = t.rows[i].cells
    cells[0].width = Cm(3.2); cells[1].width = Cm(13.6)
    r0 = cells[0].paragraphs[0].add_run(d); r0.bold = True; set_jp(r0, JP_GOTHIC); r0.font.size = Pt(10)
    r1 = cells[1].paragraphs[0].add_run(c); set_jp(r1); r1.font.size = Pt(10)
t.columns[0].width = Cm(3.2)
t.columns[1].width = Cm(13.6)

# ===== 3. 保護者の主張 =====
heading("3. 保護者の主な主張")
sub("（1）説明が二転三転していること")
body("最も問題と感じているのは、事実が二転三転している点。物事が起きた際になぜ事実をそのまま説明しないのか、なぜ隠そうとするのかという不信感がある。当初から「このような経緯で伝えてしまった。今後このようなことがないようにする」と説明していれば大きな問題にはならなかった。", bullet=True)
sub("（2）回答書の記載に関する問題点")
body("【点①：事業者の実在確認について】 回答書では「事業所の利用実態を確認したうえで担任から連絡するようにした」とあるが、当該事業者は5月1日開所であり、4月30日時点では運営実態（拠点）が存在しない。また連絡先は090で始まる携帯番号のみであり、所在地の確認にはならない。「日だまりパレット」という名称は全国に同種のものがあり、保護者は事業所名・所在地を学校に一切伝えていない。よって「実在を確認した」とは言えない。", bullet=True)
body("【点②：緊急性・安全確保について】 回答書では「学校運営を円滑に行う必要があり、安全のために伝えた」とあるが、問い合わせは5月7日の下校時刻に関するもので、約1週間の余裕があった。保護者に連絡する時間も担任が架電する時間もあり、保護者の同意取得を省略してまでの緊急性は客観的に認められない。", bullet=True)
sub("（3）漏洩した情報")
body("学校が保護者の同意なく提供した情報は、①在籍（当該校に在籍している事実）、②クラス（2年1組）、③下校時刻の3点。これらが結びつくことで、児童の所属と下校時刻が特定される。", bullet=True)
sub("（4）背景事情")
body("約5年前、親族の事業に関連して児童への殺害予告が届いた経緯があり、保護者は児童の安全に強い警戒心を持っている。下校時刻が保護者の知らないところで第三者に伝われば、誘拐等のリスクにつながりかねない。学校に提出する個人情報（勤務先・通学経路等）は事実上提出が義務的であり、それを軽んじて漏らされることは容認できない。", bullet=True)
sub("（5）事業者側の説明との食い違い")
body("回答書には「本来は保護者が確認すべき内容である」「事業者に保護者へ確認するよう伝えた」旨の記載がある。一方、事業者の説明は「保護者に架電したが連絡がつかず学校へ連絡した。副校長から『折り返す』と言われたが連絡がなかったため、再度学校へ連絡した」というもので、保護者はこちらの説明に合理性を感じている。仮に副校長が当初から明確に「保護者の同意がなければ伝えられない」と回答していれば、事業者が再度同一内容で問い合わせることは考えにくい。保護者は録音（音声データ）を保有している。", bullet=True)

# ===== 4. 校長の対応 =====
heading("4. 校長・学校の対応に関する論点")
body("校長は保護者に対し「今回の件はミスではない」と発言しており、この点が保護者に強く残っている。学校（校長・副校長）は「保護者の同意なく伝えたことは申し訳ない」と述べているが、その謝罪が保護者に十分伝わっていない。", bullet=True)
body("回答書は校長名のみで公印・署名がなく、他の学校配布文書（行事のお知らせ等）と体裁が同じであるため、スクールロイヤーより「誰が作成したか不明であり、校長が書いていないと言い逃れができる」と指摘されている。保護者は、校長が責任を持つ趣旨での署名（公印）を求めている。", bullet=True)

# ===== 5. 教育委員会の見解 =====
heading("5. 教育委員会の見解")
body("保護者の同意なく個人情報を提供した行為について、教育委員会としては「ミスではない」では片付けられない問題であると認識している。学校がどのような説明をしたとしても、結論として保護者の同意なく情報を伝えている以上、学校の対応に問題があり、指導する立場にある。", bullet=True)
body("実在確認の方法（携帯番号への折り返し、ホームページ確認）は、確実な確認とは言えず、保護者の不安はもっともであると受け止めている。", bullet=True)
body("回答書には、保護者の同意なく情報提供したことに対する謝罪の文面が欠けていると教育委員会も確認しており、本来あるべき回答として、当該事実に対する謝罪を記載すべきと考えている。", bullet=True)
body("個人情報の重みについて、校長会・副校長会の研修等の場で、事例として（個人を特定しない形で）改めて周知・指導していく。", bullet=True)

# ===== 6. 保護者の要望 =====
heading("6. 保護者の要望")
body("回答書の撤回（虚偽と受け取られる記載を含むため、一度撤回し作り直すこと）。", bullet=True)
body("保護者の同意なく情報を提供したことに対する明確な謝罪文面の作成。", bullet=True)
body("校長が責任を持つ趣旨での署名（公印）を付すこと。", bullet=True)
body("最終的に、校長・副校長、保護者（夫を含む）、教育委員会が同席し、学校として改めて謝罪する場を設けること。", bullet=True)

# ===== 7. 教育委員会の今後の対応 =====
heading("7. 今後の対応方針（教育委員会）")
body("① 教育委員会（都丸統括指導主事）より放課後等デイサービス事業者へ連絡し、4月30日当日のやり取りについて事実確認を行う（責任者＝施設長、送迎担当＝佐藤氏に確認予定）。確認の際は保護者名を出すことについて保護者の了承を得た。", bullet=True)
body("② 事実確認の結果を踏まえ、学校へ再度指導を行う。", bullet=True)
body("③ 学校に対し、保護者の同意なく情報提供したことへの謝罪、経緯・理由、再発防止を含む書面の作成を働きかける。校長の署名（公印）の要否についてはスクールロイヤーに確認のうえ対応する。", bullet=True)
body("④ 担当所管課（子育て支援課）とも相談し、過去の同種事案の有無も含めて確認する。", bullet=True)
body("⑤ 校長会・副校長会の研修において、個人情報の重要性を事例として周知・指導する。", bullet=True)
body("⑥ 上記を踏まえ、関係者同席による謝罪の場の日程調整を行う。", bullet=True)

# ===== 8. スケジュール =====
heading("8. 当面のスケジュール")
body("今週中（6月26日・金曜日まで）に、教育委員会より放課後等デイサービスへの事実確認を進め、進捗を保護者へ電話連絡する（保護者の希望は午前中）。", bullet=True)

# ===== フッター注記 =====
doc.add_paragraph().paragraph_format.space_after = Pt(6)
note = doc.add_paragraph()
rn = note.add_run("※ 本議事録は面談での発言内容を整理したものであり、事実関係（特に学校・事業者間のやり取り）については、今後の事実確認により確定される。")
set_jp(rn); rn.font.size = Pt(9); rn.italic = True
rn.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.save("/home/user/con30/議事録_20260622.docx")
print("saved")
