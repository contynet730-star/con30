# -*- coding: utf-8 -*-
"""
v11(32枚)→v12(33枚)
- 冒頭に「ペアトーク③ クイズ形式」を1枚追加（既存S4の直前）
  「主体的に学習に取り組む態度を普段何で評価していますか？」4択
  → 1分ペア相談 → 挙手で確認 → 本題（既存S4の数字あてクイズ）へ
"""
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

WATER_BLUE=RGBColor(0x5B,0xC0,0xDE); LIGHT_WATER=RGBColor(0xD6,0xEE,0xFA)
DEEP_WATER=RGBColor(0x2A,0x96,0xC2); NAVY=RGBColor(0x1F,0x3A,0x6E)
ORANGE_KW=RGBColor(0xEB,0x6C,0x15); YELLOW=RGBColor(0xFF,0xE0,0x66)
WHITE=RGBColor(0xFF,0xFF,0xFF); BLACK=RGBColor(0x00,0x00,0x00); DARK_GRAY=RGBColor(0x40,0x40,0x40)
GREEN=RGBColor(0x07,0xA9,0x73)
F_MAIN="メイリオ"

prs = Presentation('/tmp/output/講義スライド_石神井南中_20260603_v11.pptx')
SW,SH = prs.slide_width, prs.slide_height
assert len(prs.slides)==32

def blank(): return prs.slides.add_slide(prs.slide_layouts[6])

def chapter_bar(slide,text,fs=24,color=DEEP_WATER):
    bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,SW,Emu(720000))
    bar.fill.solid(); bar.fill.fore_color.rgb=color; bar.line.fill.background()
    tf=bar.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(300000); tf.margin_right=Emu(100000); tf.margin_top=Emu(20000); tf.margin_bottom=Emu(20000)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT; r=p.add_run(); r.text=text
    r.font.size=Pt(fs); r.font.bold=True; r.font.color.rgb=WHITE; r.font.name=F_MAIN

def sub_heading(slide,text,fs=28):
    bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Emu(830000),SW,Emu(700000))
    bar.fill.solid(); bar.fill.fore_color.rgb=LIGHT_WATER; bar.line.fill.background()
    tf=bar.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(300000); tf.margin_right=Emu(200000); tf.margin_top=Emu(20000); tf.margin_bottom=Emu(20000)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT; r=p.add_run(); r.text=text
    r.font.size=Pt(fs); r.font.bold=True; r.font.color.rgb=NAVY; r.font.name=F_MAIN
    ul=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Emu(1530000),SW,Emu(40000))
    ul.fill.solid(); ul.fill.fore_color.rgb=WATER_BLUE; ul.line.fill.background()

def box(slide,left,top,width,height,lines,default_size=24,default_color=BLACK,line_spacing=1.3,
        align=PP_ALIGN.LEFT,fill_color=WHITE,border_color=WATER_BLUE,border_width=3.0,
        vanchor=MSO_ANCHOR.MIDDLE,default_bold=True):
    b=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,left,top,width,height)
    b.fill.solid(); b.fill.fore_color.rgb=fill_color
    if border_color: b.line.color.rgb=border_color; b.line.width=Pt(border_width)
    else: b.line.fill.background()
    tf=b.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=Emu(140000); tf.margin_right=Emu(140000); tf.margin_top=Emu(60000); tf.margin_bottom=Emu(60000)
    tf.vertical_anchor=vanchor
    for i,item in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if isinstance(item,list):
            for j,ri in enumerate(item):
                rt,ro = ri if isinstance(ri,tuple) else (ri,{})
                if j==0: p.alignment=ro.get('align',align); p.line_spacing=ro.get('line_spacing',line_spacing)
                r=p.add_run(); r.text=rt
                r.font.size=Pt(ro.get('size',default_size)); r.font.bold=ro.get('bold',default_bold)
                r.font.color.rgb=ro.get('color',default_color); r.font.name=F_MAIN
        else:
            text,o = item if isinstance(item,tuple) else (item,{})
            p.alignment=o.get('align',align); p.line_spacing=o.get('line_spacing',line_spacing)
            r=p.add_run(); r.text=text
            r.font.size=Pt(o.get('size',default_size)); r.font.bold=o.get('bold',default_bold)
            r.font.color.rgb=o.get('color',default_color); r.font.name=F_MAIN
    return b

# ============================================================
# 冒頭クイズペアトーク（新スライド）
# ============================================================
qp = blank()
chapter_bar(qp, "【冒頭クイズ】　─　いきなりですが、ペアで", color=GREEN)
sub_heading(qp, "1分で話してください")

# 大問題
box(qp, Emu(400000), Emu(1900000), Emu(11400000), Emu(900000),
    [[("先生は、", {'size':24,'color':BLACK,'align':PP_ALIGN.CENTER})],
     [("「主体的に学習に取り組む態度」", {'size':28,'color':ORANGE_KW,'align':PP_ALIGN.CENTER}),
      ("を　普段 何で評価していますか？", {'size':24,'color':BLACK,'align':PP_ALIGN.CENTER})]],
    fill_color=LIGHT_WATER, border_width=5.0, border_color=GREEN, line_spacing=1.3)

# 4択
choices = [
    ("A", "ノート・レポートの記述内容"),
    ("B", "授業中の発言・挙手"),
    ("C", "提出物の有無・期日"),
    ("D", "教師の感覚的な印象"),
]
cw = Emu(5500000)
gx = Emu(400000)
for i,(letter, txt) in enumerate(choices):
    row, col = divmod(i, 2)
    xx = Emu(400000 + col*(5500000+400000))
    yy = Emu(2950000 + row*900000)
    # letter ball
    box(qp, xx, yy, Emu(800000), Emu(750000),
        [(letter, {'size':36,'bold':True,'color':WHITE,'align':PP_ALIGN.CENTER})],
        fill_color=DEEP_WATER, border_color=DEEP_WATER, border_width=2.0)
    box(qp, Emu(xx+900000), yy, Emu(5500000-900000), Emu(750000),
        [(txt, {'size':19,'color':NAVY,'align':PP_ALIGN.LEFT})],
        border_width=2.5)

box(qp, Emu(400000), Emu(4750000), Emu(11400000), Emu(380000),
    [("○ 一番大きい比重のものを1つ選んでください　／　○ 複数当てはまる人は「主」を選ぶ",
      {'size':16,'color':DARK_GRAY,'align':PP_ALIGN.CENTER})],
    border_color=None, fill_color=WHITE)

# タイマー＆方法
box(qp, Emu(400000), Emu(5230000), Emu(5400000), Emu(800000),
    [("⏱ 1分", {'size':36,'bold':True,'color':WHITE,'align':PP_ALIGN.CENTER})],
    fill_color=ORANGE_KW, border_color=ORANGE_KW, border_width=2.0)
box(qp, Emu(6400000), Emu(5230000), Emu(5400000), Emu(800000),
    [("隣の方とペアで", {'size':22,'bold':True,'color':NAVY,'align':PP_ALIGN.CENTER})],
    fill_color=YELLOW, border_color=ORANGE_KW, border_width=2.0)

box(qp, Emu(400000), Emu(6090000), Emu(11400000), Emu(330000),
    [("→ このあと挙手で確認します　／　答えを受けて、本日のテーマに入ります",
      {'size':14,'color':DARK_GRAY,'align':PP_ALIGN.CENTER})],
    fill_color=WHITE, border_color=None)

# ============================================================
# 並べ替え：v11の32枚 + 新規1枚 = 33枚
# 現在の順序：0..31（add した順は 0..31, 新規は 32）
# 望ましい順：S1〜S3, ★新規（冒頭クイズ）, S4〜S32
#            = [0,1,2, 32, 3,4,5,...,31]
# ============================================================
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
desired = [0,1,2, 32] + list(range(3,32))
assert sorted(desired) == list(range(33)), desired
for el in ids: sldIdLst.remove(el)
for i in desired: sldIdLst.append(ids[i])

# ============================================================
# ノート埋め込み（新冒頭クイズ）
# ============================================================
note = """【ペアトーク③ ─ 冒頭クイズ】 ─ 目安2分（相談1分＋挙手確認1分）

ウォームアップに、いきなりクイズです。

先生方ご自身に、率直にお答えいただきたいのです。
「『主体的に学習に取り組む態度』を、普段 何で評価していますか？」

A：ノート・レポートの記述内容
B：授業中の発言・挙手
C：提出物の有無・期日
D：教師の感覚的な印象

複数当てはまる方は、一番大きい比重のものを1つ。これを、隣の先生と1分間 話してみてください。

※ 1分後、講師は挙手で確認する：
　「Aの方？」「Bの方？」「Cの方？」「Dの方？」
　会場の傾向を見て、コメント：
　・Aが多ければ「皆さん意識が高い」
　・B/Cが多ければ「ここを今日は見直していきます」
　・Dも正直で良い

このあとの一言（核心の伏線）：
「実は、本日の研修の出発点が、ここなんです。
　『主体的に学習に取り組む態度』を、何で・どう見取るのか ──
　これを、皆さんと一緒に整理していくのが、今日のテーマです。」

→ そのまま次のスライド（数字あてクイズ）に入る。"""

for s in prs.slides:
    if any(sh.has_text_frame and '【冒頭クイズ】' in sh.text_frame.text for sh in s.shapes):
        s.notes_slide.notes_text_frame.text = note
        break

# フォント統一
from lxml import etree
NS_A='http://schemas.openxmlformats.org/drawingml/2006/main'
def force(run):
    rPr=run._r.find(f'{{{NS_A}}}rPr')
    if rPr is None: rPr=etree.SubElement(run._r,f'{{{NS_A}}}rPr'); run._r.insert(0,rPr)
    for t in ['latin','ea','cs']:
        e=rPr.find(f'{{{NS_A}}}{t}')
        if e is None: e=etree.SubElement(rPr,f'{{{NS_A}}}{t}')
        e.set('typeface','メイリオ')
def wf(shapes):
    n=0
    for sh in shapes:
        if sh.shape_type==MSO_SHAPE_TYPE.GROUP: n+=wf(sh.shapes)
        elif sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs: force(r); n+=1
    return n
print('font runs:', sum(wf(s.shapes) for s in prs.slides))

import os; os.makedirs('/tmp/output',exist_ok=True)
OUT='/tmp/output/講義スライド_石神井南中_20260603_v12.pptx'
prs.save(OUT)
print('saved:', OUT, 'slides:', len(prs.slides))
