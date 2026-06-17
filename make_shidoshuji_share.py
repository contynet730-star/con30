# -*- coding: utf-8 -*-
"""
【指導主事 共有用】学校公開日（土曜授業）の日程に関する保護者対応の考え方・想定問答 を生成。

目的：保護者対応（今井様の件）で上司から
  「土曜日の背景・変えられないこと等に触れないと納得されない」と指摘された点を踏まえ、
  同種のお問い合わせに各指導主事が同じ考え方で対応できるよう、説明の骨子と想定問答を共有する。
出典：添付3資料（PDF＝土曜授業の在り方／PPT＝令和8年度の土曜日授業の在り方／当初回答案docx）
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21.0)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

MINCHO = "ＭＳ 明朝"
GOTHIC = "ＭＳ ゴシック"


def set_font(run, size, bold=False, color=None, name=MINCHO):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def para(text="", size=10, bold=False, align=None, name=MINCHO,
         before=0, after=3, line=15, indent=None, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = Pt(line)
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Cm(indent)
    if text:
        r = p.add_run(text)
        set_font(r, size, bold=bold, color=color, name=name)
    return p


def heading(text):
    para(text, size=11, bold=True, name=GOTHIC, before=7, after=3, line=16,
         color=(0, 51, 102))


def bullet(text, indent=0.5):
    para("・" + text, size=10, indent=indent, after=2, line=15)


# ── タイトル ─────────────────────────────────────────
para("【指導主事 共有用】", size=10.5, bold=True, name=GOTHIC,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para("学校公開日（土曜授業）の日程に関する保護者対応の考え方・想定問答",
     size=13, bold=True, name=GOTHIC, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("～「背景」と「変えられないこと」を踏まえた説明のために～",
     size=10, name=GOTHIC, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
para("令和８年６月　　教育振興部　教育指導課",
     size=9.5, align=WD_ALIGN_PARAGRAPH.RIGHT, after=4)

# 0 趣旨
heading("０　共有の趣旨")
para("　保護者（今井様）からの学校公開日（土曜授業）の日程に関するお問い合わせに対し、当初の"
     "回答案は「重要な視点として受け止めるが、趣旨をご理解いただきたい」という結論のみで、"
     "①なぜ土曜授業を行うのか（背景）、②なぜ日程を個別に変えられないのか（理由）が示されて"
     "おらず、相手の納得が得られにくいとの指摘があった。同種のお問い合わせは各指導主事にも"
     "寄せられうるため、説明の考え方を共有し、対応の標準化を図るもの。", after=3)

# 1 骨子
heading("１　説明の骨子（３点セットで構成する）")
bullet("（１）背景　……　制度の経緯と目的を示し、思いつきの運用ではないことを伝える。")
bullet("（２）変えられないこと　……　個別変更が難しい理由を率直に示す。")
bullet("（３）変えられること　……　令和８年度からの柔軟化で前向きに応え、否定で終わらせない。")

# 2 背景
heading("２　「背景」の説明材料（出典：別添「土曜授業の在り方」）")
bullet("完全学校週５日制（平成14年度～）の下、学校・家庭・地域が連携し社会全体で子供を育てる。")
bullet("文部科学省：平成25年11月、学校教育法施行規則を改正＝設置者の判断で土曜授業が可能と明確化。")
bullet("東京都教育委員会：平成20年12月、開かれた学校づくりの観点から土曜授業実施の留意点を通知。")
bullet("練馬区：平成24年度から、振替休業日を設定しない土曜授業を全小・中学校で実施。")
bullet("ねらい＝①必要な授業時数の確保、②保護者・地域への教育活動の公開（理解と信頼の向上）。", indent=1.0)

# 3 制度的根拠（管理運営規則）
heading("３　制度的根拠：練馬区立学校の管理運営に関する規則（土曜授業関連）")
bullet("休業日：土曜日・日曜日等は規則上「休業日」（＝本来は授業を行わない日）と定められている。")
bullet("休業日の授業：休業日に授業を行う（又は授業日に休業する）ときは、校長が教育委員会の許可を"
       "受ける必要がある（運動会等、年間行事計画に基づく恒常的行事は事前の届出で足りる）。"
       "→土曜授業は区の枠組みの中で計画的に実施するもので、個別の裁量で随時動かせるものではない。")
bullet("回数の上限：土曜日における教育課程に位置付けられた授業の実施は「各月２回を上限」とされている。")
para("　※「各月２回」はあくまで制度上の上限。現行の年４回程度は、働き方改革による負担軽減のため"
     "上限より少なく設定した政策判断であり、「上限まで増やせる」との誤解を招かないよう、回数の"
     "説明では上限を前面に出さない。", size=9.5, after=3, line=15)
para("　※条番号・正確な文言は例規集（練馬区立学校の管理運営に関する規則）で確認のこと。",
     size=9.5, after=3, line=15)

# 4 変えられないこと
heading("４　「変えられないこと」とその理由（出典：在り方／R8資料／管理運営規則）")
bullet("全区共通・計画的な日程設定：目的が「多くの保護者・地域に開く」ことであり、地域行事や区の"
       "各種事業を計画する際の指標にもなっている。→特定のご家庭ごとの日程設定は制度趣旨・学校運営"
       "上できない。")
bullet("実施回数は増やせない：働き方改革により児童・生徒・教職員の負担軽減を図る趣旨で、令和６年度"
       "に年８回→年４回程度へ縮減済み。")
bullet("制度の枠組み：上記３のとおり、休業日である土曜に授業を行うには教育委員会の許可を要し、"
       "区全体の計画の中で実施している。個別のご家庭に合わせた変更にはなじまない。")

# 5 変えられること
heading("５　「変えられること」＝令和８年度からの見直し（出典：R8資料）")
bullet("「原則第２土曜日」→第２土曜日以外の土曜日にも実施可（各学校からの改善要望に対応）。")
bullet("年４回→年３～４回（必ず１時間以上の授業公開を実施）。")
bullet("振替休業日：設定可（給食あり・５～６時間授業が条件）※給食提供は保健給食課と調整中。")
para("　→　保護者の「日程が固定されていることによる負担」に直接応える材料。"
     "まずは『原則第２土曜の柔軟化』を中心にご案内する。", size=10, after=3, line=15)

# 6 想定問答
heading("６　想定問答（Q&A）")
qa = [
    ("Q1　なぜ土曜日に行うのか。平日ではだめか。",
     "A1　開かれた学校づくりと授業時数の確保が目的。保護者・地域が参観しやすい休業日として土曜を"
     "活用しており、国・都も土曜授業を制度上位置付けている。"),
    ("Q2　第２土曜は都合が悪い。日にちを変えてほしい。",
     "A2　個別のご家庭ごとの設定は制度趣旨・運営上難しいが、令和８年度から第２土曜以外も実施可に"
     "見直した。学校の年間予定をご確認のうえ学校へご相談を。平日の学校公開・授業参観もご案内する。"),
    ("Q3　回数を増やせないのか。",
     "A3　働き方改革による負担軽減の趣旨で年８回→年４回程度に縮減した経緯があり、回数増は困難。"),
    ("Q4　給食付き・午後までの公開はできないか。",
     "A4　令和８年度から振替休業日の設定（給食あり・５～６時間授業）が可能となる方向で調整中。"
     "実施は学校・区の決定による。※確定状況を確認のうえ回答する。"),
    ("Q5　土曜に授業を行ってよい制度的根拠は。回数の決まりはあるのか。",
     "A5　区立学校の管理運営に関する規則上、土曜日は休業日だが、教育委員会の許可の下で休業日にも"
     "授業を実施できる。土曜の教育課程上の授業は各月２回が上限。本区はこの枠組みの中で計画的に"
     "実施している（現行は働き方改革により年４回程度）。"),
]
for q, a in qa:
    para(q, size=10, bold=True, name=GOTHIC, before=3, after=1, line=15)
    para("　" + a, size=10, after=2, line=15, indent=0.3)

# 7 留意点
heading("７　対応上の留意点")
bullet("「各月２回」は制度上の上限であり、現行の運用回数（年４回程度）とは異なる。"
       "回数を問われても、上限までの増加を約束しない。")
bullet("令和８年度の運用変更（特に給食・振替の取扱い）の確定状況を、回答前に必ず確認する"
       "（保健給食課との調整事項）。")
bullet("個別の確約は避け、「学校の年間予定の確認」「学校への相談」を案内する導線にとどめる。")
bullet("回答は『背景→変えられないこと→令和８年度の改善』の順で、否定で終わらず前向きな材料で締める。")

para("", after=2)
para("【参考】出典：練馬区立学校の管理運営に関する規則（練馬区例規集）、別添「土曜授業の在り方」"
    "及び「令和８年度の土曜日授業の在り方」資料。", size=9, after=4, line=14)
para("担当：教育振興部　教育指導課（電話　５９８４－５７５９）", size=9.5,
     align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)

doc.save("/home/user/con30/【指導主事共有用】土曜授業の日程に関する対応の考え方.docx")
print("Saved: 【指導主事共有用】土曜授業の日程に関する対応の考え方.docx")
