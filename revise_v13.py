# -*- coding: utf-8 -*-
"""
v12(33枚)→v13: ペアトーク3枚に公的資料の根拠を明示
- S4 冒頭クイズ：中教審報告(H31.1)「学習評価の現状の課題」
- S13 ペアトーク①：指導要領 解説 総則編 第3章第3節1
- S28 ペアトーク②：国研「指導と評価の一体化のための学習評価
                       に関する参考資料」（各教科）
各スライドの「→ このあと…」予告ボックスに、本問いの根拠を併記。
発表者ノートにも「この問いの根拠」セクションを追加。
"""
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE_TYPE

DARK_GRAY = RGBColor(0x40,0x40,0x40)
NAVY = RGBColor(0x1F,0x3A,0x6E)
ORANGE_KW = RGBColor(0xEB,0x6C,0x15)
GREEN = RGBColor(0x07,0xA9,0x73)
F_MAIN = "メイリオ"

prs = Presentation('/tmp/output/講義スライド_石神井南中_20260603_v12.pptx')

# ============================================================
# 各ペアトークスライド下部「→ このあと…」を「予告＋根拠」の二行構成に書き換え
# ============================================================
# (slide_index, 予告テキスト, 根拠テキスト)
updates = [
    (3,
     "→ このあと挙手で確認します　／　答えを受けて、本日のテーマに入ります",
     "※本問いの根拠：中央教育審議会『児童生徒の学習評価の在り方について（報告）』（平成31年1月）「学習評価の現状の課題」"),
    (12,
     "→ このあと、「主体的に学習に取り組む態度」の評価をどう見取るか、お話します",
     "※本問いの根拠：中学校学習指導要領 解説 総則編 第3章第3節1「主体的・対話的で深い学びの実現に向けた授業改善」"),
    (27,
     "→ ここで描いた「Ｂの姿」を、このあとのグループ協議で持ち寄ります",
     "※本作業の根拠：国立教育政策研究所「指導と評価の一体化のための学習評価に関する参考資料（中学校）」各教科版"),
]

for slide_idx, expected_preview, source_text in updates:
    s = prs.slides[slide_idx]
    target = None
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("→"):
            target = sh
            break
    if target is None:
        print(f'WARN: S{slide_idx+1} 予告ボックスが見つかりません')
        continue

    # 既存テキストを書き換え：予告は短縮し、根拠を2段落目で追加
    tf = target.text_frame
    # 既存の全paragraphsを削除して再構築
    from lxml import etree
    NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    # txBody内の全p要素を削除
    txBody = tf._txBody
    for p_el in txBody.findall(f'{{{NS_A}}}p'):
        txBody.remove(p_el)
    # 2段落を新規追加
    # 1段落目：予告
    p1_el = etree.SubElement(txBody, f'{{{NS_A}}}p')
    pPr1 = etree.SubElement(p1_el, f'{{{NS_A}}}pPr')
    pPr1.set('algn','ctr')
    r1_el = etree.SubElement(p1_el, f'{{{NS_A}}}r')
    rPr1 = etree.SubElement(r1_el, f'{{{NS_A}}}rPr')
    rPr1.set('sz','1300'); rPr1.set('b','0')
    fill1 = etree.SubElement(rPr1, f'{{{NS_A}}}solidFill')
    clr1 = etree.SubElement(fill1, f'{{{NS_A}}}srgbClr')
    clr1.set('val','404040')
    for tag in ['latin','ea','cs']:
        f = etree.SubElement(rPr1, f'{{{NS_A}}}{tag}')
        f.set('typeface','メイリオ')
    t1 = etree.SubElement(r1_el, f'{{{NS_A}}}t')
    t1.text = expected_preview
    # 2段落目：根拠
    p2_el = etree.SubElement(txBody, f'{{{NS_A}}}p')
    pPr2 = etree.SubElement(p2_el, f'{{{NS_A}}}pPr')
    pPr2.set('algn','ctr')
    r2_el = etree.SubElement(p2_el, f'{{{NS_A}}}r')
    rPr2 = etree.SubElement(r2_el, f'{{{NS_A}}}rPr')
    rPr2.set('sz','1100'); rPr2.set('b','1')
    fill2 = etree.SubElement(rPr2, f'{{{NS_A}}}solidFill')
    clr2 = etree.SubElement(fill2, f'{{{NS_A}}}srgbClr')
    clr2.set('val','076F40')  # 落ち着いた緑
    for tag in ['latin','ea','cs']:
        f = etree.SubElement(rPr2, f'{{{NS_A}}}{tag}')
        f.set('typeface','メイリオ')
    t2 = etree.SubElement(r2_el, f'{{{NS_A}}}t')
    t2.text = source_text

# ============================================================
# 発表者ノートに「この問いの根拠」を追記
# ============================================================
note_addenda = {
    3: """

─────────────────────
【この問いの根拠】

本問いは、中央教育審議会「児童生徒の学習評価の在り方について（報告）」
（平成31年1月21日）が示した「学習評価の現状の課題」を踏まえて設定。

報告では、観点別学習状況の評価について、特に「関心・意欲・態度」
（現「主体的に学習に取り組む態度」の前身）が、ともすると
「ノート提出頻度」「挙手回数」など、表面的・形式的な評価に
なりがちであると指摘されている。

【4択選択肢A〜Dとの対応】
A：ノート・レポートの記述内容
  → 記述の「内容」を見るなら正しい方向（本研修で扱う）
B：授業中の発言・挙手
  → 報告書が「形式的評価」の典型として問題視
C：提出物の有無・期日
  → 報告書が指摘する「真面目さ」評価への陥穽
D：教師の感覚的な印象
  → 妥当性・信頼性ともに低い（S24で扱う）

→ どの答えが多くても、講師は否定せず受け止め、「今日はそこを
   一緒に整理していきます」と本題への伏線として活用する。""",

    12: """

─────────────────────
【この問いの根拠】

本問いは、中学校学習指導要領（平成29年告示）解説 総則編
第3章第3節1「主体的・対話的で深い学びの実現に向けた授業改善」
の記述に基づく。

解説では、3つの学びについて以下のように例示されている：
- 主体的な学び：見通し／粘り強さ／振り返り
- 対話的な学び：互いの考えの比較／共に考えを創る
- 深い学び：知識・技能の活用／概念化／問い続ける

これらは「目指す生徒の姿」として例示されており、本研修のS9-S11
（各6要素）の出典でもある。

→ ペアトーク①は、抽象的な理論を「自分の授業の具体的場面」に
   照らし戻す、解説書の意図に沿った振り返り作業である。

【講師の振る舞い】
ペア対話中は教室を巡回し、出てきた具体例を1-2件メモする。
発表後に「先ほど○○先生のグループから『◯◯』というお話が
聞こえました。これはまさに『△△な学び』の好例ですね」と
還元すると、根拠が体感される。""",

    27: """

─────────────────────
【この作業の根拠】

本作業は、国立教育政策研究所「『指導と評価の一体化』のための
学習評価に関する参考資料（中学校）」（令和2年3月）の各教科版が
示す「内容のまとまりごとの評価規準」を、自分の単元に具体化する
作業である。

国研の参考資料では、各教科で「Bと判断する姿」が例示されている
ものの、それは「内容のまとまり」レベル。実際の授業では、これを
個々の単元・場面に降ろす必要がある。この降ろし作業こそが、
教師の専門性であり、教科会で揃える価値の源泉でもある。

→ つまり、この3分の作業は「国研資料を自分の教科で読み解く」
   ことに他ならない。

【講師の振る舞い】
ペア対話中の巡回で、特に具体的に書けている例を1-2件メモし、
ペア終了後に「○○先生がこんなふうに書かれていました」と
紹介すると、他の先生方の参考になる。
（事前に許可を取るのが丁寧）

【グループ協議への接続】
ここで描いた「Bの姿」を、各教科会で持ち寄って統合する。
それが、このあとの研究協議会のメイン作業になる。""",
}

for idx, addendum in note_addenda.items():
    s = prs.slides[idx]
    if s.has_notes_slide:
        existing = s.notes_slide.notes_text_frame.text
        s.notes_slide.notes_text_frame.text = existing + addendum

# ============================================================
# フォント統一
# ============================================================
from lxml import etree
NS_A='http://schemas.openxmlformats.org/drawingml/2006/main'
def force(run):
    rPr=run._r.find(f'{{{NS_A}}}rPr')
    if rPr is None: rPr=etree.SubElement(run._r,f'{{{NS_A}}}rPr'); run._r.insert(0,rPr)
    for t in ['latin','ea','cs']:
        e=rPr.find(f'{{{NS_A}}}{t}')
        if e is None: e=etree.SubElement(rPr,f'{{{NS_A}}}{t}')
        e.set('typeface','メイリオ')
def wf(shapes):
    n=0
    for sh in shapes:
        if sh.shape_type==MSO_SHAPE_TYPE.GROUP: n+=wf(sh.shapes)
        elif sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs: force(r); n+=1
    return n
print('font runs:', sum(wf(s.shapes) for s in prs.slides))

import os; os.makedirs('/tmp/output',exist_ok=True)
OUT='/tmp/output/講義スライド_石神井南中_20260603_v13.pptx'
prs.save(OUT)
print('saved:', OUT, 'slides:', len(prs.slides))
